import sys, codecs
with codecs.open(r'C:\GOOGLE ANGET\三合一單自動產生器\main.py', 'r', 'utf-8') as f:
    content = f.read()

# 1. Headers
content = content.replace(
    '            (8, "修正到廠時間"),\n            (9, "單列清空")',
    '            (8, "修正到廠時間"),\n            (9, "採購單號"),\n            (10, "單列清空")'
)

# 2. UI Generation
old_ui = """            # Col 9: 單列清空按鈕
            btn_clear_row = tk.Button(
                self.scrollable_frame, 
                text="清空", 
                command=lambda r=row_idx-1: self.clear_single_row(r), 
                font=("Microsoft JhengHei", 9, "bold"), 
                bg="#FFEBEE", 
                fg="#C62828", 
                cursor="hand2", 
                width=6,
                pady=1
            )
            btn_clear_row.grid(row=row_grid_idx, column=9, padx=4, pady=2)"""

new_ui = """            # Col 9: 採購單號
            po_var = tk.StringVar(value="")
            po_entry = tk.Entry(self.scrollable_frame, textvariable=po_var, width=18, font=("Arial", 10), fg="#333")
            po_entry.grid(row=row_grid_idx, column=9, padx=2, pady=2, sticky="ew")

            # Col 10: 單列清空按鈕
            btn_clear_row = tk.Button(
                self.scrollable_frame, 
                text="清空", 
                command=lambda r=row_idx-1: self.clear_single_row(r), 
                font=("Microsoft JhengHei", 9, "bold"), 
                bg="#FFEBEE", 
                fg="#C62828", 
                cursor="hand2", 
                width=6,
                pady=1
            )
            btn_clear_row.grid(row=row_grid_idx, column=10, padx=4, pady=2)"""
content = content.replace(old_ui, new_ui)

# 3. Add to entries dict
content = content.replace(
    '                "mod_time_var": mod_time_var\n            })',
    '                "mod_time_var": mod_time_var,\n                "po_var": po_var\n            })'
)

# 4. clear_all_rows
content = content.replace(
    '                entry["mod_time_var"].set("")',
    '                entry["mod_time_var"].set("")\n                if "po_var" in entry: entry["po_var"].set("")'
)

# 5. clear_single_row
content = content.replace(
    '            entry["mod_time_var"].set("")',
    '            entry["mod_time_var"].set("")\n            if "po_var" in entry: entry["po_var"].set("")'
)

# 6. CSV extraction PO logic
content = content.replace(
    'batch_col, loc_col, date_col, tank_col, time_col, mod_time_col = -1, -1, -1, -1, -1, -1',
    'batch_col, loc_col, date_col, tank_col, time_col, mod_time_col, po_col = -1, -1, -1, -1, -1, -1, -1'
)
content = content.replace(
    'if mod_time_col == -1 and "修正" in v and ("時間" in v or "TIME" in v): mod_time_col = c_idx',
    'if mod_time_col == -1 and "修正" in v and ("時間" in v or "TIME" in v): mod_time_col = c_idx\n                                if po_col == -1 and any(k in v for k in ["採購單", "PO"]): po_col = c_idx'
)
content = content.replace(
    'mt_val = normalize_time_str(get_c(mod_time_col))',
    'mt_val = normalize_time_str(get_c(mod_time_col))\n                            po_val = str(get_c(po_col) or "").strip()'
)
content = content.replace(
    '                                    "mod_time": mt_val\n                                })',
    '                                    "mod_time": mt_val,\n                                    "po": po_val\n                                })'
)

# 7. Excel extraction PO logic
content = content.replace(
    'if mod_time_idx == -1 and "修正" in h_str and ("時間" in h_str or "TIME" in h_str): mod_time_idx = c_idx',
    'if mod_time_idx == -1 and "修正" in h_str and ("時間" in h_str or "TIME" in h_str): mod_time_idx = c_idx\n                                if po_idx == -1 and any(k in h_str for k in ["採購單", "PO"]): po_idx = c_idx'
)
content = content.replace(
    'batch_idx = loc_idx = date_idx = tank_idx = time_idx = mod_time_idx = -1',
    'batch_idx = loc_idx = date_idx = tank_idx = time_idx = mod_time_idx = po_idx = -1'
)
content = content.replace(
    'mt_val = normalize_time_str(get_val(mod_time_idx))',
    'mt_val = normalize_time_str(get_val(mod_time_idx))\n                            po_val = str(get_val(po_idx) or "").strip()'
)
content = content.replace(
    '                                "mod_time": mt_val\n                            })',
    '                                "mod_time": mt_val,\n                                "po": po_val\n                            })'
)

# 8. load_coa_forms logic for PO
content = content.replace(
    '                default_date = row["date_var"].get().strip()\n                \n                date_str = default_date',
    '                default_date = row["date_var"].get().strip()\n                \n                po_val = row.get("po_var", tk.StringVar()).get().strip()[:10]\n                \n                date_str = default_date'
)
content = content.replace(
    '                    ws[\'B7\'] = default_phase\n                    ws[\'B11\'] = default_date',
    '                    ws[\'B7\'] = default_phase\n                    ws[\'B11\'] = default_date\n                    ws[\'B12\'] = po_val'
)
content = content.replace(
    '                    if len(reader) > 6: reader[6][1] = default_phase\n                    if len(reader) > 10: reader[10][1] = default_date\n                    \n                    with open(new_file_path, \'w\', encoding=\'utf-8-sig\', newline=\'\') as f:',
    '                    if len(reader) > 6: reader[6][1] = default_phase\n                    if len(reader) > 10: reader[10][1] = default_date\n                    if len(reader) > 11: reader[11][1] = po_val\n                    \n                    with open(new_file_path, \'w\', encoding=\'utf-8-sig\', newline=\'\') as f:'
)

# 9. load_coa_forms Lorry extraction
old_lorry = """                default_phase = "F" + loc_str if loc_str else ""
                default_date = row["date_var"].get().strip()
                
                po_val = row.get("po_var", tk.StringVar()).get().strip()[:10]
                
                date_str = default_date"""
new_lorry = """                default_phase = "F" + loc_str if loc_str else ""
                default_date = row["date_var"].get().strip()
                
                tsmc_fab, fab_phase, deliv_date = lorry_data.get(matched_batch, (default_fab, default_phase, default_date))
                if not deliv_date: deliv_date = default_date
                
                po_val = row.get("po_var", tk.StringVar()).get().strip()[:10]
                
                date_str = deliv_date
                factory_code = tsmc_fab"""
content = content.replace(old_lorry, new_lorry)
content = content.replace("ws['B6'] = default_fab", "ws['B6'] = tsmc_fab")
content = content.replace("ws['B7'] = default_phase", "ws['B7'] = fab_phase")
content = content.replace("ws['B11'] = default_date", "ws['B11'] = deliv_date")
content = content.replace("reader[5][1] = default_fab", "reader[5][1] = tsmc_fab")
content = content.replace("reader[6][1] = default_phase", "reader[6][1] = fab_phase")
content = content.replace("reader[10][1] = default_date", "reader[10][1] = deliv_date")

old_lorry_read = """        error_msgs = []

        for file_path in file_paths:"""
new_lorry_read = """        error_msgs = []
        
        lorry_data = {}
        if getattr(self, "imported_lorry_files", None) and self.imported_lorry_files:
            try:
                import openpyxl
                self.update()
                wb_l = openpyxl.load_workbook(self.imported_lorry_files[0], data_only=True, read_only=True)
                ws_l = wb_l.active
                headers = []
                batch_idx = fab_idx = phase_idx = date_idx = -1
                for idx, row_data in enumerate(ws_l.iter_rows(values_only=True)):
                    if idx % 100 == 0: self.update()
                    if not headers and idx < 10:
                        row_strs = [str(c or "").strip() for c in row_data]
                        if "FinalBatchID" in row_strs:
                            headers = row_strs
                            batch_idx = headers.index("FinalBatchID")
                            fab_idx = headers.index("TSMCFab") if "TSMCFab" in headers else -1
                            phase_idx = headers.index("FabPhase") if "FabPhase" in headers else -1
                            date_idx = headers.index("DeliveryDate") if "DeliveryDate" in headers else -1
                        continue
                    if headers and batch_idx != -1 and row_data[batch_idx]:
                        b = str(row_data[batch_idx]).strip().upper()
                        f_v = str(row_data[fab_idx]).strip() if fab_idx != -1 and row_data[fab_idx] is not None else ""
                        p_v = str(row_data[phase_idx]).strip() if phase_idx != -1 and row_data[phase_idx] is not None else ""
                        d_v = row_data[date_idx] if date_idx != -1 else ""
                        from datetime import datetime
                        if isinstance(d_v, datetime): d_v = f"{d_v.year}/{d_v.month:02d}/{d_v.day:02d}"
                        else: d_v = str(d_v or "").strip()
                        lorry_data[b] = (f_v, p_v, d_v)
                wb_l.close()
            except Exception as e:
                print(f"Read lorry error: {e}")

        for file_path in file_paths:"""
content = content.replace(old_lorry_read, new_lorry_read)

# 10. Also missing `po_val` mapping in generate_files setup
content = content.replace(
    '                    if rec.get("mod_time"): row_e["mod_time_var"].set(rec["mod_time"])',
    '                    if rec.get("mod_time"): row_e["mod_time_var"].set(rec["mod_time"])\n                    if rec.get("po"): row_e.get("po_var", tk.StringVar()).set(rec["po"])'
)

# 11. Vertical CSV PO logic
content = content.replace(
    '                        if batch_val:',
    '                        po_val = ""\n                        if batch_val:'
)
content = content.replace(
    '                                "mod_time": "",\n                                "tank": get_tank_from_batch(batch_val)\n                            })',
    '                                "mod_time": "",\n                                "tank": get_tank_from_batch(batch_val),\n                                "po": po_val\n                            })'
)

with codecs.open(r'C:\GOOGLE ANGET\三合一單自動產生器\main.py', 'w', 'utf-8') as f:
    f.write(content)
print("PO logic and Lorry data mapping restored.")
