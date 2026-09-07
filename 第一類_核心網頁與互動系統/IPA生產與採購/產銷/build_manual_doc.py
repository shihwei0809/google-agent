import os
try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("缺少 python-docx 套件，請執行 pip install python-docx")
    exit(0)

def generate_manual():
    doc = Document()
    
    # 標題
    title = doc.add_heading('勝一化工 - 產銷計畫表 Web 系統操作手冊', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('1. 系統緣由與設計目標', level=1)
    doc.add_paragraph('本系統專為勝一化工產銷排程設計，將 Excel「20260824產銷計畫.xlsx」全面轉化為高互動性之獨立 Web HTML 系統。')
    doc.add_paragraph('具備獨立運行目錄（不與 IPA 核心程式混雜）、總表優先視圖、多月份排程切換與嚴謹4個月滾動結存連動計算。')

    doc.add_heading('2. 核心功能規格', level=1)
    doc.add_paragraph('• 📊 分頁架構：第一頁固定為「產銷總表」，依序為「精餾產銷」、「初餾產銷」、「調配產銷」、「🧪 耗料 (BOM 換算)」、「⚙️ 生產線別參照」。')
    doc.add_paragraph('• 📅 月份排程切換填報：上方提供 [8月排程] [9月排程] [10月排程] [11月排程] 快速按鈕，切換時即時帶出該月份 3 次生產期間（線別、日期、數量 MT），無需水平冗長捲動。')
    doc.add_paragraph('• 🔄 嚴謹4個月滾動結存公式：')
    doc.add_paragraph('    1. 8/31 結存 = 08/24期初 + 8月生產總量(數量1+2+3) - 08/24~08/31銷')
    doc.add_paragraph('    2. 9/30 結存 = 8/31結存 + 9月生產總量(數量1+2+3) - 09/30銷')
    doc.add_paragraph('    3. 10/31 結存 = 9/30結存 + 10月生產總量(數量1+2+3) - 10/31銷')
    doc.add_paragraph('    4. 11/30 結存 = 10/31結存 + 11月生產總量(數量1+2+3) - 11/30銷')
    doc.add_paragraph('• 👁️ 隱藏無生產/需求產品：一鍵切換過濾無庫存、無銷售且無排產之品項。')
    doc.add_paragraph('• 📥 匯出 CSV 報表：一鍵匯出包含 4 個月各期生產與連續結存之完整 CSV 報表。')

    doc.add_heading('3. 快速啟動方式', level=1)
    doc.add_paragraph('• 方式一：直接雙擊「點我開啟產銷計畫系統.bat」，自動以預設瀏覽器開啟。')
    doc.add_paragraph('• 方式二：直接以瀏覽器開啟 src/index.html 即可離線單機執行。')

    filename = os.path.join(os.path.dirname(__file__), '產銷計畫系統_操作手冊.docx')
    doc.save(filename)
    print(f"手冊已成功生成：{filename}")

if __name__ == "__main__":
    generate_manual()
