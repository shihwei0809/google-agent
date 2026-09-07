import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_manual():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("【系統操作手冊】勝一化工 - 產銷計畫 Web 系統 (PWA 獨立 App 版)")
    run_title.font.name = "微軟正黑體"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(14, 116, 144)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("版本：v7.0 (PWA 雙軌版) | 適用設備：Windows 電腦 / Android 手機 / iPhone / 平板")
    run_sub.font.name = "微軟正黑體"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph("―" * 50)
    
    doc.add_heading("一、 系統簡介與 4 個月滾動生產公式", level=1)
    doc.add_paragraph("本系統為勝一化工產銷計畫表之現代化 PWA 系統，直接對齊 Excel 產銷計畫，具備 4 個月連續滾動連動公式：")
    
    formulas = [
        ("1. 8/31 預估結存", "08/24 期初庫存 + 8月生產總量 (數量1+2+3) - 08/24~08/31 銷"),
        ("2. 9/30 預估結存", "8/31 庫存 + 9月生產總量 (數量1+2+3) - 09/30 銷"),
        ("3. 10/31 預估結存", "9/30 庫存 + 10月生產總量 (數量1+2+3) - 10/31 銷"),
        ("4. 11/30 預估結存", "10/31 庫存 + 11月生產總量 (數量1+2+3) - 11/30 銷")
    ]
    for f_title, f_calc in formulas:
        p = doc.add_paragraph()
        r = p.add_run(f"◆ {f_title} = ")
        r.bold = True
        r.font.color.rgb = RGBColor(14, 116, 144)
        p.add_run(f_calc)

    doc.add_heading("二、 核心亮點功能", level=1)
    highlights = [
        ("產銷總表置頂", "首頁直接呈現產銷總表，並提供精餾、初餾、調配、耗料BOM、生產線別參照快速頁籤。"),
        ("月份切換排程", "點選「8月 / 9月 / 10月 / 11月 排程」按鈕，立即切換當月 3 次排程（線別、日期區間、數量），即時連動最右側 4 個月結存。"),
        ("隱藏無生產/需求產品", "勾選「👁️ 隱藏無生產/需求產品」，畫面立即過濾無需求品項，聚焦於當月實際運作項目。"),
        ("雲端 GAS 串接", "內建 GAS Web App 快速啟動與雲端同步設定，方便與既有 Google Apps Script 資料庫對接。")
    ]
    for h_title, h_desc in highlights:
        p = doc.add_paragraph()
        r = p.add_run(f"★ {h_title}：")
        r.bold = True
        p.add_run(h_desc)

    doc.add_heading("三、 PWA 獨立 App 安裝與離線使用指引", level=1)
    steps = [
        ("Windows 電腦端 (Chrome/Edge)", "雙擊「啟動PWA本機測試.bat」開啟網頁，點擊頁面頂部「立即安裝」或網址列右側「安裝勝一產銷計畫」圖示，即可將系統安裝為桌面獨立 App。"),
        ("Android 手機 / 平板", "手機連線本機伺服器之區域網 IP (例如 http://192.168.x.x:8086)，點擊頂部「立即安裝」或瀏覽器選單「新增至主畫面」。"),
        ("iPhone / iPad (Safari)", "在 Safari 中開啟網址，點擊底部「分享」按鈕，選擇「加入主畫面」，即可如同原生 App 般在主畫面開啟使用。")
    ]
    for s_title, s_desc in steps:
        p = doc.add_paragraph()
        r = p.add_run(f"【{s_title}】\n")
        r.bold = True
        r.font.color.rgb = RGBColor(14, 116, 144)
        p.add_run(s_desc)

    base = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(base, "勝一產銷計畫系統_操作手冊.docx")
    doc.save(docx_path)
    print(f"Word 手冊已成功生成: {docx_path}")

    try:
        import win32com.client
        word = win32com.client.Dispatch('Word.Application')
        doc_obj = word.Documents.Open(docx_path)
        pdf_path = os.path.join(base, "勝一產銷計畫系統_操作手冊.pdf")
        doc_obj.SaveAs(pdf_path, FileFormat=17)
        doc_obj.Close()
        word.Quit()
        print(f"PDF 手冊已成功生成: {pdf_path}")
    except Exception as e:
        print(f"PDF 轉換略過或由系統自動產生: {e}")

if __name__ == '__main__':
    create_manual()
