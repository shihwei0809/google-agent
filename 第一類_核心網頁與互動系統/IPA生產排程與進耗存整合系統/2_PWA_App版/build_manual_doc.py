import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_manual():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading("IPA 生產排程與進耗存整合系統", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("PWA 雙軌獨立 App 與雲端整合操作手冊 (v3.1)")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.color.rgb = RGBColor(79, 70, 229)
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.bold = True
    
    doc.add_paragraph()
    
    # Section 1
    doc.add_heading("一、 系統簡介與特色", level=1)
    p = doc.add_paragraph(
        "本系統專為 IPA 彰濱一廠（5 槽共用原料池）與二廠（獨立原料槽）打造，"
        "整合產銷排產、機台流速試算、儲槽每日進耗存預估與跨月連續繼承。"
    )
    doc.add_paragraph("核心亮點：", style='List Bullet')
    doc.add_paragraph("【PWA 雙軌架構】：支援電腦與手機離線安裝為獨立 App，免開啟瀏覽器網址列。", style='List Bullet')
    doc.add_paragraph("【最新產銷日期起算】：產銷排產模式下自動以最新填寫之日期（如 10/12）為起算點，無需手動清空歷史舊日期。", style='List Bullet')
    doc.add_paragraph("【雙核心動態試算】：同時支援全月滿載模式與產銷實體原料重量起算模式。", style='List Bullet')
    doc.add_paragraph("【雙雲端儲存支援】：支援 Google 試算表 (GAS) 與 Cloudflare D1 雲端資料庫雙向儲存。", style='List Bullet')
    
    # Section 2
    doc.add_heading("二、 產銷排產模式使用說明", level=1)
    doc.add_paragraph("1. 點擊右側功能面板切換至「📦 產銷當日原料庫存」分頁。")
    doc.add_paragraph("2. 選擇廠區（一廠共用池或二廠獨立原料槽）。")
    doc.add_paragraph("3. 在月曆日期格子中輸入實體原料重量（例如在 12 日輸入 438.5 噸）。")
    doc.add_paragraph("4. 系統自動在 12 日加上「起算」紫色標籤，下方進耗存表格自動由 12 日起算並裁剪至月底。")
    doc.add_paragraph("5. 若前面有舊日期（如 5 日），系統會自動標示為「(舊)」，絕不干擾 12 日的計算。")
    
    # Section 3
    doc.add_heading("三、 PWA 桌面與手機安裝步驟", level=1)
    doc.add_paragraph("1. 使用 Chrome 或 Edge 瀏覽器開啟系統網址（https://ipa-pwa-app.pages.dev）。")
    doc.add_paragraph("2. 畫面上方會自動彈出「📲 安裝 PWA 桌面/手機版應用程式！」提示條。")
    doc.add_paragraph("3. 點擊「立即安裝」按鈕，即可將系統安裝至桌面或手機主畫面。")
    doc.add_paragraph("4. 本機運行時亦可直接執行「啟動PWA本機測試.bat」進行本機離線伺服器測試。")
    
    output_docx = "IPA生產排程與進耗存整合系統_操作手冊.docx"
    doc.save(output_docx)
    print(f"操作手冊已成功產出：{output_docx}")

if __name__ == "__main__":
    generate_manual()
