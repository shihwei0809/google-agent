import codecs
path = r'd:\GOOGLE ANGET\勝一三合一單網頁架機伺服器\server.py'
with codecs.open(path, 'r', 'utf-8-sig') as f:
    lines = f.readlines()

new_lines = []
in_ext = False
for i, line in enumerate(lines):
    # 1. Fix _extract_batch_from_coa_bytes
    if 'def _extract_batch_from_coa_bytes(content, ext):' in line:
        in_ext = True
        new_lines.append("""def _extract_batch_from_coa_bytes(content, ext):
    import io
    import csv
    try:
        f = io.StringIO(content.decode('utf-8-sig', errors='ignore'))
        reader = csv.reader(f)
        for row in reader:
            if len(row) >= 2 and 'RawLotId' in str(row[0]):
                return str(row[1]).strip()
    except:
        pass
    try:
        f = io.StringIO(content.decode('cp950', errors='ignore'))
        reader = csv.reader(f)
        for row in reader:
            if len(row) >= 2 and 'RawLotId' in str(row[0]):
                return str(row[1]).strip()
    except:
        pass
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(min_row=1, max_row=50, min_col=1, max_col=2):
                if row[0].value and 'RawLotId' in str(row[0].value) and row[1].value:
                    return str(row[1].value).strip()
    except:
        pass
    return None
""")
        continue
    if in_ext:
        if 'return None' in line:
            in_ext = False
        continue
    
    # 2. Fix the loop logic
    if 'if not matched_batch:' in line and 'continue' in lines[i+1]:
        indent = ' ' * (len(line) - len(line.lstrip()))
        new_lines.append(indent + 'if not matched_batch:\n')
        new_lines.append(indent + '    found_batch = _extract_batch_from_coa_bytes(coa_file["content"], ext)\n')
        new_lines.append(indent + '    if found_batch:\n')
        new_lines.append(indent + '        for b in valid_records:\n')
        new_lines.append(indent + '            if b == found_batch or b in found_batch or found_batch in b:\n')
        new_lines.append(indent + '                matched_batch = b\n')
        new_lines.append(indent + '                break\n')
        continue

    # 3. Fix import os issue by replacing local import os with prefix generation
    if 'idx = base_name.upper().find(matched_batch)' in line:
        indent = ' ' * (len(line) - len(line.lstrip()))
        new_lines.append(indent + 'idx = base_name.upper().find(matched_batch)\n')
        new_lines.append(indent + 'if idx == -1:\n')
        new_lines.append(indent + '    import os as _os\n')
        new_lines.append(indent + '    name_no_ext, ext_part = _os.path.splitext(base_name)\n')
        new_lines.append(indent + '    prefix = name_no_ext + "_"\n')
        new_lines.append(indent + '    suffix = ext_part\n')
        new_lines.append(indent + '    idx = len(prefix)\n')
        new_lines.append(indent + '    base_name = prefix + matched_batch + suffix\n')
        new_lines.append(indent + 'else:\n')
        new_lines.append(indent + '    prefix = base_name[:idx]\n')
        new_lines.append(indent + '    suffix = base_name[idx + len(matched_batch):]\n')
        continue
    
    if 'prefix = base_name[:idx]' in line and 'suffix = base_name[idx + len(matched_batch):]' in lines[i+1]:
        continue
    if 'suffix = base_name[idx + len(matched_batch):]' in line and 'prefix = base_name[:idx]' in lines[i-1]:
        continue

    new_lines.append(line)

with codecs.open(path, 'w', 'utf-8-sig') as f:
    f.writelines(new_lines)
