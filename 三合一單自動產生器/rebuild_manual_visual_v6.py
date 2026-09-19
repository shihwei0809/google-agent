from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
ASSET = ROOT / "manual_assets"
VISUALS = ASSET / "v5_visuals"
OUT = ROOT / "三合一單與單列生產履歷 Chemical Lorry 自動產生器 圖文操作手冊 V7.docx"
FONT = "Microsoft JhengHei"


def font(run, size=10.5, bold=False, color="222222"):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def page_field(p):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def paragraph(doc, text, size=10.5, bold=False, color="222222", align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    font(r, size, bold, color)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    font(r, 16, True, "000000")


def figure(doc, file, caption, width=15.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(file), width=Cm(width))
    paragraph(doc, caption, 9, True, "666666", WD_ALIGN_PARAGRAPH.CENTER)


def bullets(doc, lines):
    for line in lines:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        font(p.add_run(line), 10.5)


def step_page(doc, heading_text, intro, visual, caption, notes):
    heading(doc, heading_text)
    paragraph(doc, intro)
    figure(doc, VISUALS / visual, caption)
    bullets(doc, notes)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(fp.add_run("三合一單與單列生產履歷 圖文操作手冊 V7  |  第 "), 8.5, False, "666666")
    page_field(fp)
    font(fp.add_run(" 頁"), 8.5, False, "666666")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(8)
    font(p.add_run("三合一單與單列生產履歷\nChemical Lorry 自動產生器"), 24, True, "000000")
    paragraph(doc, "圖文操作手冊 V7", 13, True, "2E7D32", WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "以目前實際啟動的程式畫面編寫。每一步先看紅框位置，再依短句說明操作；不需要閱讀長篇文字即可完成日常出貨作業。", 11, False, "333333", WD_ALIGN_PARAGRAPH.CENTER)
    figure(doc, ASSET / "current_main_ui.png", "總覽：先由上到下確認系統狀態、資料載入、資料列與最下方的產生按鈕", 17.6)
    doc.add_page_break()

    step_page(doc, "步驟 1 開始前確認系統檔案", "啟動程式後，先看最上方系統狀態。範本與地點代號對照表必須顯示已找到；若本次需要單列生產履歷，先選取正確的 Chemical Lorry 來源檔。", "01_status.png", "圖 1 紅框為重新讀取對照表與載入生產履歷來源檔的位置", [
        "地點代號對照表剛修改時，按「重新載入對照表」後才開始輸入。",
        "只產生三合一單時，不必載入生產履歷來源檔。",
        "來源檔載入成功後，狀態列會顯示檔名並自動勾選生產履歷輸出。",
    ])
    step_page(doc, "步驟 2 選擇資料來源", "工具列的三個按鈕用途不同：排程資料使用 Excel 匯入；單列生產履歷使用 Chemical Lorry 來源檔；已存在的 COA Excel 或 CSV 表單使用 COA 載入。", "02_import.png", "圖 2 依本次工作選擇正確載入方式", [
        "要建立三合一單時，通常先按「從 Excel 匯入排程」。",
        "「載入 COA 表單」是更新既有 COA 檔案，不是上傳圖片截圖。",
        "匯入後仍需在主列表檢查批號、地點與出貨日期。",
    ])
    step_page(doc, "步驟 3 設定批次出貨日期", "當多筆資料共用同一出貨日期時，先輸入批次日期，再按「套用至全列」。系統只會套用到已有批號的資料列。", "03_batch.png", "圖 3 使用全選、日期與套用按鈕加速批次作業", [
        "日期建議使用 YYYY/MM/DD，後續最容易核對輸出資料夾。",
        "按「帶入今天日期」只會補入目前空白的日期。",
        "左側勾選框決定該列是否參與產生。",
    ])
    step_page(doc, "步驟 4 填寫與核對資料列", "每一筆要產生的資料都必須保留勾選，填入 10 碼批號與地點代號。槽號及長代號由程式自動帶入；出貨日期可在個別列上調整。", "04_rows.png", "圖 4 紅框標示日常作業最常使用的資料欄位", [
        "批號必須剛好 10 碼，否則無法產生。",
        "地點必須存在於地點代號對照表；找不到時先更新對照表。",
        "單列有誤時按右側「清空」，不要誤按右上角的全部清除。",
    ])
    step_page(doc, "步驟 5 勾選要輸出的報表", "三合一單 Excel 是程式目前必要的輸出項目。若已載入 Chemical Lorry 來源檔，可同時勾選單列生產履歷，程式將為每筆對應批號輸出單列檔案。", "05_reports.png", "圖 5 產生前確認兩個報表選項", [
        "三合一單會帶入槽號、批號、長代號並重新產生 QR Code。",
        "生產履歷只會產出可在來源檔找到相同批號的資料。",
        "不確定來源檔是否正確時，先只產生一筆測試確認。",
    ])
    step_page(doc, "步驟 6 開始產生與驗收", "完成資料核對後，按下最下方綠色按鈕。處理期間請等待完成訊息，勿重複點擊。程式完成後會開啟依日期、地點與槽號整理的輸出資料夾。", "06_generate.png", "圖 6 綠色按鈕會開始產生所有已勾選資料列的 Excel 檔", [
        "完成視窗會顯示成功份數及資料夾位置。",
        "三合一單會放在「三合一單輸出_YYYYMMDD」下的地點與槽號子資料夾。",
        "交付前開啟抽查：批號、槽號、地點長代號、出貨日期與生產履歷列。",
    ])
    heading(doc, "COA 表單更新與常見狀況")
    paragraph(doc, "若要更新 COA，先在主列表完成並勾選批號，然後選取「載入 COA 表單」。程式會依檔名或表單內批號配對並另存輸出檔，不直接覆蓋原始 COA。")
    bullets(doc, [
        "COA 無法處理：先確認主列表已勾選對應批號，且 COA 檔名或內容可辨識批號。",
        "地點找不到：更新地點代號對照表後按「重新載入對照表」。",
        "生產履歷未產出：確認已載入來源檔，且來源檔內存在相同批號。",
        "產生失敗：依錯誤訊息回到指定資料列，優先檢查批號長度、地點與日期。",
    ])
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
