from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "三合一單與單列生產履歷 Chemical Lorry 自動產生器 操作手冊 V5.docx"
UI_IMAGE = ROOT / "manual_assets" / "current_main_ui.png"
FONT = "Microsoft JhengHei"
NAVY = "17365D"
BLUE = "1F4E78"
GREEN = "2E7D32"
ORANGE = "C55A11"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def border(cell, color="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def cell_text(cell, text, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text)
    set_font(r, size, bold, color)
    border(cell)


def add_page_field(paragraph):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, 15 if level == 1 else 12, True, NAVY if level == 1 else BLUE)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, 10.5, True, NAVY)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r, 10.3)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Cm(widths[i])
        shade(cell, NAVY)
        cell_text(cell, title, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
    for row_no, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Cm(widths[i])
            if row_no % 2:
                shade(cells[i], "F7FAFC")
            cell_text(cells[i], value, i == 0, None, WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_step(doc, number, title, action, check):
    add_heading(doc, f"{number}. {title}", 2)
    add_body(doc, action, "操作：")
    add_body(doc, check, "確認：")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles["Normal"].font.size = Pt(10.5)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("三合一單與單列生產履歷自動產生器 操作手冊 V5  |  第 ")
    set_font(fr, 8.5, False, "666666")
    add_page_field(fp)
    fr = fp.add_run(" 頁")
    set_font(fr, 8.5, False, "666666")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(45)
    title.paragraph_format.space_after = Pt(10)
    r = title.add_run("三合一單與單列生產履歷\nChemical Lorry 自動產生器")
    set_font(r, 24, True, NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(32)
    r = subtitle.add_run("現場作業操作手冊  V5")
    set_font(r, 13, True, GREEN)
    add_body(doc, "本手冊依 2026 年 9 月 19 日實際啟動的主程式畫面與目前 main.py 行為重新編寫。適用於以排程資料建立三合一單，並選擇性輸出單列生產履歷與更新 COA 表單的作業人員。")
    add_body(doc, "使用前請先確認範本、地點代號對照表與必要的生產履歷檔案均由管理人員提供。手冊中的檔名、資料夾與欄位名稱以目前程式畫面為準。")
    doc.add_page_break()

    add_heading(doc, "1 先看懂主畫面", 1)
    add_body(doc, "啟動後先從上到下確認系統狀態，再進行資料匯入。主要工作區由系統狀態、資料載入工具列、批次設定、報表勾選與資料列組成。")
    if UI_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(UI_IMAGE), width=Cm(16.8))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run("圖 1 目前版本主畫面")
        set_font(r, 9, True, "666666")
    add_table(doc, ["區域", "用途", "開始作業前要確認的事"], [
        ("系統狀態", "確認範本、地點代號對照表與 Chemical Lorry 來源檔。", "前兩項必須顯示已找到；若要產生生產履歷，先載入來源檔。"),
        ("資料載入", "匯入 Excel 排程、載入生產履歷、更新 COA 表單。", "三種功能對應不同工作，不要混用為同一匯入動作。"),
        ("批次設定", "勾選資料列、設定共用出貨日期。", "只會套用到已輸入批號的資料列。"),
        ("資料列表", "逐筆檢查批號、地點、出貨日期、採購單號。", "批號必須剛好 10 碼；地點必須存在對照表。"),
    ], [2.8, 7.1, 7.1])
    doc.add_page_break()

    add_heading(doc, "2 啟動與開工檢查", 1)
    add_step(doc, "2.1", "啟動程式", "在程式資料夾中雙擊「啟動_三合一單產生器.bat」。啟動器會檢查 Python 與 openpyxl、qrcode、Pillow、pytesseract 等套件後開啟主畫面。", "看到「三合一單自動產生器」主視窗即可；不要關閉顯示環境檢查的命令視窗，關閉主程式後它才會結束。")
    add_step(doc, "2.2", "確認必要檔案", "查看最上方「系統狀態」。程式需要「台積電槽車barcode三合一單-範本.xlsx」及「地點代號對照表.xlsx」。", "兩項均顯示綠色「已找到」。若地點對照表剛被修改，按「重新載入對照表」再開始。")
    add_step(doc, "2.3", "選擇生產履歷來源", "若這次需要單列 Chemical Lorry，按橘色「選擇生產履歷檔」或工具列「載入生產履歷 (Chemical_Lorry)」，選取正確的 Excel 檔。", "狀態列會顯示已就緒的檔名，且「產生單列生產履歷 Excel」會自動勾選。")
    add_body(doc, "注意：若僅要產生三合一單，不必載入 Chemical Lorry；但「產生三合一單 Excel」必須保持勾選，程式目前不支援單獨產生其他報表。", "注意：")

    add_heading(doc, "3 建立或匯入排程資料", 1)
    add_step(doc, "3.1", "從 Excel 匯入排程", "按藍色「從 Excel 匯入排程」，可選取 Excel 或 CSV 檔案。依彈出視窗的日期與筆數篩選，確認預覽後匯入。", "資料會填入主畫面列表；每筆仍請核對批號、地點與出貨日期。")
    add_step(doc, "3.2", "手動輸入或貼上", "在「批號」欄輸入 10 碼批號，於「地點」欄輸入如 15P5 的代號。可從 Excel 複製資料後在列表貼上，仍須逐筆檢視自動帶出的內容。", "槽號與長代號為自動欄位；若未帶出，先檢查批號格式與地點對照表。")
    add_step(doc, "3.3", "批次補日期與調整列數", "用「帶入今天日期」補入空白日期；或在「批次出貨日期」填日期後按「套用至全列」。資料超過畫面列數時按「新增 10 列」。", "每一筆要產生的資料列左側皆保持勾選，且出貨日期正確。")
    add_step(doc, "3.4", "清除錯誤資料", "按資料列最右側「清空」只清該筆；按右上「清除全部資料」會要求確認後清除全部已輸入欄位。", "重新輸入後，槽號與長代號應隨批號、地點變更。")
    doc.add_page_break()

    add_heading(doc, "4 產生三合一單與單列生產履歷", 1)
    add_step(doc, "4.1", "選擇本次輸出", "在「欲產生的報表勾選」區確認「產生三合一單 Excel」已勾選；需要 Chemical Lorry 時，同時勾選其選項並確保已載入來源檔。", "不要勾選空白或不需要的資料列，避免產出多餘檔案。")
    add_step(doc, "4.2", "執行產生", "按最下方綠色「開始批次產生 Excel 報表」。程式會檢查必填欄位、批號長度與地點代號，處理時請等待完成訊息。", "完成後程式會顯示成功份數與資料夾位置，並開啟輸出資料夾。")
    add_table(doc, ["檢查項目", "程式規則", "處理方式"], [
        ("批號", "必須剛好 10 碼。", "修正資料列後再次產生。"),
        ("地點", "必須存在於地點代號對照表。", "修正代號，或更新對照表後按重新載入。"),
        ("日期", "可接受 YYYY/MM/DD、YYYY-MM-DD、YYYY.MM.DD、YYYYMMDD 等格式；空白時以當日分組。", "建議統一使用 YYYY/MM/DD，方便核對輸出資料夾。"),
        ("生產履歷", "要產生單列檔案時，來源檔必須有可對應的批號列。", "確認來源檔批號與主畫面批號一致。"),
    ], [2.8, 7.0, 7.2])

    doc.add_page_break()
    add_heading(doc, "5 輸出位置與檔案結構", 1)
    add_body(doc, "程式會依出貨日期分組，在程式資料夾內建立「三合一單輸出_YYYYMMDD」資料夾；其下再以「MMDD 地點 槽號」建立子資料夾。例如，同日不同地點或槽號會分開存放。")
    add_table(doc, ["輸出項目", "檔案位置與命名", "驗收重點"], [
        ("三合一單", "三合一單輸出_YYYYMMDD\\MMDD 地點 槽號\\YYYY.M.D. 槽號 地點台積電槽車barcode三合一單.xlsx", "確認槽號、批號、地點長代號、QR Code 與版面。"),
        ("單列生產履歷", "與對應三合一單放在同一子資料夾，檔名保留來源檔名稱並加上日期、槽號與地點。", "確認輸出為對應批號的單列資料。"),
        ("工作快取", "last_generated_session.json 會記錄本次有效資料。", "僅供程式記錄，不是交付檔案。"),
    ], [2.8, 8.1, 6.1])
    doc.add_page_break()

    add_heading(doc, "6 載入 COA 表單並更新欄位", 1)
    add_body(doc, "紫色「載入 COA 表單」處理的是既有的 COA Excel 或 CSV 表單，並非舊手冊所稱的「上傳 COA 截圖」。使用前必須先在主列表填好、勾選要對應的批號。")
    add_step(doc, "6.1", "準備對應資料", "在主列表中輸入並勾選批號、地點與出貨日期；如要帶入生產履歷資料，先載入 Chemical Lorry 來源檔。", "COA 檔名或表單內容必須能辨識出對應批號，否則程式會列入錯誤紀錄。")
    add_step(doc, "6.2", "選取 COA 檔案", "按「載入 COA 表單」，可多選 .xlsx、.xls 或 .csv。程式會以檔名優先比對批號，檔名沒有時再嘗試從表單讀取。", "完成訊息顯示成功處理份數；有未對應檔案時依錯誤紀錄逐一處理。")
    add_step(doc, "6.3", "確認寫入結果", "程式會另存到對應的輸出日期與地點槽號資料夾，不直接覆蓋原檔。沒有載入生產履歷時，會寫入 COA 的 B6；有對應生產履歷時，可帶入 B6、B7、B11，採購單號前 10 碼可寫入 B12。", "開啟另存檔確認欄位內容、日期與檔名正確，再交付或上傳。")
    add_body(doc, "注意：COA 功能的欄位內容取決於已選取資料列與生產履歷是否成功比對；對不到批號時不要手動覆寫原始 COA，請先保留錯誤訊息再查來源資料。", "注意：")

    doc.add_page_break()
    add_heading(doc, "7 常見狀況", 1)
    add_table(doc, ["現象", "可能原因", "建議處理"], [
        ("地點代號找不到", "對照表未更新或代號輸入錯誤。", "更新地點代號對照表後按「重新載入對照表」，再重試。"),
        ("無法產生", "批號不是 10 碼、地點或必填值不完整。", "依錯誤視窗指定的項次修正；只留下確定要輸出的勾選列。"),
        ("Chemical Lorry 未產出", "未載入來源檔或來源檔沒有相同批號。", "重新選擇正確的來源檔，確認批號格式與大小寫。"),
        ("COA 沒有被處理", "主列表未勾選批號，或 COA 內容/檔名無法比對。", "先填寫並勾選資料列，再以可辨識批號的檔名重新載入。"),
        ("畫面與舊手冊不同", "舊版手冊描述了已變動的截圖 OCR 按鈕與舊流程。", "以本 V5 手冊和目前主畫面的按鈕文字為準。"),
    ], [3.1, 6.2, 7.7])

    doc.add_page_break()
    add_heading(doc, "8 每次交付前的最後檢查", 1)
    add_bullets(doc, [
        "確認每份三合一單的批號、槽號、地點長代號與出貨日期。",
        "確認產出資料夾日期、地點與槽號符合本次出貨資料。",
        "若有 Chemical Lorry，確認該檔為正確批號的單列版本。",
        "若有 COA 表單，確認另存檔欄位及檔名已更新，原始檔仍可保留。",
        "完成後再關閉主程式；如有錯誤，保留完成訊息與來源檔供追查。",
    ])
    add_body(doc, "版本說明：V5 重新依目前主程式的按鈕、資料列與輸出規則編寫，取代先前與現況不一致的截圖操作說明。")
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
