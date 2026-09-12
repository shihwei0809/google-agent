import re, codecs
with codecs.open(r'C:\GOOGLE ANGET\三合一單自動產生器\main.py', 'r', 'utf-8') as f:
    content = f.read()

lorry_read_code = r'''        success_count = 0
        error_msgs = []
        
        lorry_data = {}
        if getattr(self, "imported_lorry_files", None) and self.imported_lorry_files:
            try:
                import openpyxl
                wb_l = openpyxl.load_workbook(self.imported_lorry_files[0], data_only=True)
                ws_l = wb_l.active
                headers = []
                batch_idx = fab_idx = phase_idx = date_idx = -1
                for idx, row_data in enumerate(ws_l.iter_rows(values_only=True)):
                    if not headers and idx < 10:
                        row_strs = [str(c or "").strip() for c in row_data]
                        if "FinalBatchID" in row_strs:
                            headers = row_strs
                            batch_idx = headers.index("FinalBatchID")
                            fab_idx = headers.index("TSMCFab") if "TSMCFab" in headers else -1
                            phase_idx = headers.index("FabPhase") if "FabPhase" in headers else -1
                            date_idx = headers.index("DeliveryDate") if "DeliveryDate" in headers else -1
                        continue
                    if headers and batch_idx != -1 and row_data[batch_idx]:
                        b = str(row_data[batch_idx]).strip().upper()
                        f_v = str(row_data[fab_idx]).strip() if fab_idx != -1 and row_data[fab_idx] is not None else ""
                        p_v = str(row_data[phase_idx]).strip() if phase_idx != -1 and row_data[phase_idx] is not None else ""
                        d_v = row_data[date_idx] if date_idx != -1 else ""
                        from datetime import datetime
                        if isinstance(d_v, datetime): d_v = f"{d_v.year}/{d_v.month:02d}/{d_v.day:02d}"
                        else: d_v = str(d_v or "").strip()
                        lorry_data[b] = (f_v, p_v, d_v)
                wb_l.close()
            except Exception as e:
                print(f"Read lorry error: {e}")
'''
content = re.sub(r'        success_count = 0\s+error_msgs = \[\]', lorry_read_code.strip('\r\n'), content)

new_logic = r'''                row = valid_batches[matched_batch]
                loc_str = row["loc_var"].get().strip()
                long_code = row["long_code_var"].get().strip()
                
                if long_code and len(long_code) >= 5:
                    default_fab = long_code[1:5]
                else:
                    default_fab = loc_str[1:5] if len(loc_str) >= 5 else loc_str
                
                default_phase = "F" + loc_str if loc_str else ""
                default_date = row["date_var"].get().strip()
                
                tsmc_fab, fab_phase, deliv_date = lorry_data.get(matched_batch, (default_fab, default_phase, default_date))
                if not deliv_date: deliv_date = default_date
                
                po_val = row.get("po_var", tk.StringVar()).get().strip()[:10]
                
                date_str = deliv_date
                formatted_date = date_str.replace("/", "").replace("-", "")
                
                factory_code = tsmc_fab'''

content = re.sub(r'                row = valid_batches\[matched_batch\]\s+loc_str = row\[\"loc_var\"\].get\(\).strip\(\)\s+factory_code = loc_str\[1:5\] if len\(loc_str\) >= 5 else loc_str\s+date_str = row\[\"date_var\"\].get\(\).strip\(\)\s+formatted_date = date_str.replace\(\"/\", \"\"\).replace\(\"-\", \"\"\)', new_logic, content)

new_write_excel = r'''                    ws['B6'] = tsmc_fab
                    ws['B7'] = fab_phase
                    ws['B11'] = deliv_date
                    ws['B12'] = po_val'''
content = re.sub(r'                    ws\[\'B6\'\] = factory_code\s+ws\[\'B17\'\] = date_str', new_write_excel, content)

new_write_csv = r'''                    reader[5][1] = tsmc_fab
                    if len(reader) > 6: reader[6][1] = fab_phase
                    if len(reader) > 10: reader[10][1] = deliv_date
                    if len(reader) > 11: reader[11][1] = po_val'''
content = re.sub(r'                    reader\[5\]\[1\] = factory_code\s+reader\[16\]\[1\] = date_str', new_write_csv, content)

with codecs.open(r'C:\GOOGLE ANGET\三合一單自動產生器\main.py', 'w', 'utf-8') as f:
    f.write(content)
print("Patch 3 applied.")
