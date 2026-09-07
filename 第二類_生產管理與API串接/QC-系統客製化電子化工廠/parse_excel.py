# -*- coding: utf-8 -*-
"""
T100 槽車排程解析腳本 (支援出貨與進料雙區塊解析)
專門鎖定「明天進出貨排程報表(槽車)」分頁
自動識別：
  1. 上半部「出貨」區塊 (出貨通知單)
  2. 下半部「進料」區塊 (入庫單)
"""
import openpyxl
import sys
import re
import json
import os

sys.stdout.reconfigure(encoding='utf-8')

def parse_t100_tank_schedule(excel_path):
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"找不到檔案: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # 專門鎖定槽車分頁
    target_sheet_name = None
    for name in wb.sheetnames:
        if '槽車' in name:
            target_sheet_name = name
            break
    
    if not target_sheet_name:
        target_sheet_name = wb.sheetnames[0]
        print(f"⚠️ 未找到含'槽車'之工作表，預設使用: {target_sheet_name}")
    else:
        print(f"✅ 成功鎖定目標工作表: {target_sheet_name}")

    sheet = wb[target_sheet_name]

    current_mode = None
    headers = {}
    orders = []

    for r in range(1, sheet.max_row + 1):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
        row_str = ' '.join([str(v) for v in row_vals if v is not None])
        
        # 檢查是否為出貨標題列
        if '出貨通知單' in row_str:
            current_mode = '出貨'
            headers = {str(sheet.cell(row=r, column=c).value).strip(): c for c in range(1, sheet.max_column + 1) if sheet.cell(row=r, column=c).value}
            print(f"Row {r:02d}: 切換至 [出貨] 區塊")
            continue
            
        # 檢查是否為進料/入庫標題列
        if '入庫單' in row_str:
            current_mode = '進料'
            headers = {str(sheet.cell(row=r, column=c).value).strip(): c for c in range(1, sheet.max_column + 1) if sheet.cell(row=r, column=c).value}
            print(f"Row {r:02d}: 切換至 [進料] 區塊")
            continue
            
        if not current_mode:
            continue
            
        def gv(col_name):
            c = headers.get(col_name)
            return sheet.cell(row=r, column=c).value if c else None

        if current_mode == '出貨':
            doc_no = gv('出貨通知單')
            prod = gv('品名')
            if not doc_no or not prod:
                continue
            tank = gv('儲位名稱') or ''
            spec = gv('規格') or ''
            customer = gv('交易對象簡稱') or ''
            truck = gv('出通單單頭.車牌號碼') or gv('車牌號碼') or gv('車牌號碼(磅)') or ''
            qty = gv('換算數量') or ''
            note = gv('注意事項') or ''
            time_str = gv('預計出貨時間') or ''
            
            container = str(truck).strip() if truck else ''
            if not container and note:
                m = re.search(r'槽號[:：]?\s*([A-Za-z0-9\-]+)', str(note))
                if m:
                    container = m.group(1)
                else:
                    m2 = re.search(r'提貨槽車\s*([A-Za-z0-9\-]+)', str(note))
                    if m2:
                        container = m2.group(1)
                        
            grade = '工業級'
            if 'UPS' in str(prod).upper():
                grade = 'UPS'
            elif 'IF' in str(prod).upper():
                grade = 'IF'
                
            orders.append({
                'doc_no': str(doc_no).strip(),
                'time': str(time_str).strip(),
                'flowType': '出貨',
                'grade': grade,
                'productName': str(prod).strip(),
                'tankNo': str(tank).strip(),
                'customer': str(customer).strip(),
                'container': container if container else str(customer).strip(),
                'quantity': f"{qty} KG" if qty else str(spec).strip(),
                'note': str(note).strip()
            })
            
        elif current_mode == '進料':
            doc_no = gv('入庫單')
            prod = gv('品名')
            if not doc_no or not prod:
                continue
            tank = gv('槽别') or gv('儲位') or ''
            spec = gv('規格') or ''
            vendor = gv('供應商簡稱') or ''
            origin = gv('出貨廠別(廠商)') or ''
            truck = gv('車牌號碼') or ''
            box = gv('櫃號') or ''
            qty = gv('預計進貨數量(KG)') or ''
            note = gv('備註') or ''
            time_str = gv('預計進貨時間') or ''
            
            container = str(truck).strip()
            if box and str(box).strip():
                container = f"{container} (櫃:{box})" if container else str(box).strip()
                
            cust_str = f"{vendor} ({origin})" if vendor and origin else (origin or vendor or '')
            
            grade = '工業級'
            if 'UPS' in str(prod).upper():
                grade = 'UPS'
            elif 'IF' in str(prod).upper():
                grade = 'IF'
                
            orders.append({
                'doc_no': str(doc_no).strip(),
                'time': str(time_str).strip(),
                'flowType': '進料',
                'grade': grade,
                'productName': str(prod).strip(),
                'tankNo': str(tank).strip(),
                'customer': cust_str.strip(),
                'container': container if container else cust_str.strip(),
                'quantity': f"{qty} KG" if qty else str(spec).strip(),
                'note': str(note).strip()
            })

    return orders

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))
    excel_file = os.path.join(base_dir, "明天進出貨排程報表(全部)20x3A00.xlsx")
    
    parsed = parse_t100_tank_schedule(excel_file)
    print(f"\n✅ 成功完整解析 {len(parsed)} 筆槽車排程（包含出貨與進料）：")
    for i, o in enumerate(parsed, 1):
        print(f"  [{i:02d}] [{o['flowType']} {o['time']:<5}] 單號: {o['doc_no']:<20} | 品名: {o['productName']:<8} | 槽號: {o['tankNo']:<6} | 車/櫃: {o['container']:<18} | 客戶: {o['customer']:<20} | 數量: {o['quantity']}")
        
    out_json = os.path.join(base_dir, "t100_tomorrow_schedule.json")
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(parsed, f, ensure_ascii=False, indent=2)
    print(f"\n已輸出至: {out_json}")
