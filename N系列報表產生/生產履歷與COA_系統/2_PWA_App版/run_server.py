import os
import re
import socket
import pandas as pd
import openpyxl
from datetime import datetime
from fastapi import FastAPI, HTTPException
from fastapi.responses import HTMLResponse, JSONResponse
from pydantic import BaseModel
import uvicorn

app = FastAPI()

def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"

def find_available_port(start_port: int, max_attempts: int = 50) -> int:
    for port in range(start_port, start_port + max_attempts):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            try:
                s.bind(("0.0.0.0", port))
                return port
            except OSError:
                continue
    return start_port

def get_base_dir():
    # 假設在 D:\GOOGLE ANGET\N系列報表產生\生產履歷與COA_系統\2_PWA_App版
    return os.path.dirname(os.path.dirname(os.path.dirname(__file__)))

def auto_detect_files():
    base_dir = get_base_dir()
    map_path = None
    ship_path = None
    if os.path.exists(base_dir):
        for f in os.listdir(base_dir):
            if not f.endswith('.xlsx'): continue
            if '地點代號對照表' in f and not map_path:
                map_path = os.path.join(base_dir, f)
            if '出貨表' in f and not ship_path:
                ship_path = os.path.join(base_dir, f)
    return map_path, ship_path

@app.get("/api/shipping_list")
def get_shipping_list():
    map_path, ship_path = auto_detect_files()
    if not ship_path:
        return JSONResponse({"error": "找不到出貨表"}, status_code=400)
    
    try:
        # 讀取所有 sheet 並合併
        xl = pd.ExcelFile(ship_path)
        ship_data_list = []
        
        for sheet_name in xl.sheet_names:
            ship_df = pd.read_excel(ship_path, sheet_name=sheet_name, header=1)
            cols = list(ship_df.columns)
            col_idx = {}
            for i, c in enumerate(cols):
                c_str = str(c).strip()
                if '出貨' in c_str: col_idx['出貨'] = i
                elif '到貨' in c_str and '地點' not in c_str: col_idx['到貨'] = i
                elif '批號' in c_str: col_idx['批號'] = i
                elif '桶數' in c_str: col_idx['桶數'] = i
                elif '總重(L)' in c_str or '總重' in c_str: col_idx['總重'] = i
                elif '到貨地點' in c_str: col_idx['到貨地點'] = i
                elif '採購單號' in c_str: col_idx['採購單號'] = i
                
            for _, row in ship_df.iterrows():
                if '批號' not in col_idx: continue
                batch_no = str(row.iloc[col_idx['批號']]).strip()
                if batch_no == 'nan' or not batch_no: continue
                
                raw_loc = str(row.iloc[col_idx.get('到貨地點', 5)]).strip()
                clean_loc = raw_loc.replace('廠', '') 
                
                po_no_raw = str(row.iloc[col_idx.get('採購單號', 6)]).strip()
                po_no = po_no_raw.split('/')[0] if '/' in po_no_raw else po_no_raw
                
                delivery_date = row.iloc[col_idx.get('到貨', 1)]
                if isinstance(delivery_date, pd.Timestamp):
                    delivery_date = delivery_date.strftime('%Y/%m/%d')
                else:
                    delivery_date = str(delivery_date).split(' ')[0]
                    
                qty = str(row.iloc[col_idx.get('總重', 4)]).strip()
                
                ship_data_list.append({
                    "batch_no": batch_no,
                    "location": raw_loc,
                    "clean_loc": clean_loc,
                    "delivery_date": delivery_date,
                    "po_no": po_no,
                    "qty": qty,
                    "sheet": sheet_name
                })
        
        # 反轉順序讓最新的在上面
        ship_data_list.reverse()
        return {"data": ship_data_list, "ship_path": ship_path}
    except Exception as e:
        import traceback
        return JSONResponse({"error": str(e), "trace": traceback.format_exc()}, status_code=500)

class ProcessRequest(BaseModel):
    source_folder: str
    selected_batches: list[dict] # [{"batch_no": "...", "clean_loc": "..."}]

@app.post("/api/process_local")
def process_local(req: ProcessRequest):
    map_path, _ = auto_detect_files()
    if not map_path:
        return JSONResponse({"error": "找不到地點代號對照表"}, status_code=400)
        
    src_dir = req.source_folder
    if not os.path.exists(src_dir):
        return JSONResponse({"error": f"資料夾不存在: {src_dir}"}, status_code=400)

    try:
        mapping_df = pd.read_excel(map_path)
        loc_map = {}
        for _, row in mapping_df.iterrows():
            if pd.isna(row.iloc[0]): continue
            fac = str(row.iloc[0]).replace('廠', '').strip()
            code = str(row.iloc[1]).strip()
            loc_map[fac] = code
            
        selected_map = {(item['batch_no'], item['clean_loc']): item for item in req.selected_batches}
        
        out_dir = os.path.join(src_dir, f"處理後_{datetime.now().strftime('%Y%m%d_%H%M%S')}")
        os.makedirs(out_dir, exist_ok=True)
        
        files = [f for f in os.listdir(src_dir) if os.path.isfile(os.path.join(src_dir, f))]
        success_count = 0
        processed_files = []
        
        for filename in files:
            if not (filename.lower().endswith('.xlsx') or filename.lower().endswith('.csv')):
                continue
                
            filepath = os.path.join(src_dir, filename)
            out_filepath = os.path.join(out_dir, filename)
            
            m = re.search(r'([0-9]{4,5}[A-Z][0-9]{4,5})\s+.*?\s+([0-9A-Z]+)\.(xlsx|csv)', filename, re.IGNORECASE)
            if not m:
                parts = filename.replace('.xlsx', '').replace('.csv', '').split(' ')
                if len(parts) >= 2:
                    batch_no = parts[-3] if '-' in parts[-3] else parts[-3].split('-')[-1]
                    loc_part = parts[-1]
                else:
                    continue
            else:
                batch_no = m.group(1)
                loc_part = m.group(2)
                
            data = selected_map.get((batch_no, loc_part))
            if not data:
                continue

            code = loc_map.get(loc_part, '')
            delivery_date = data.get('delivery_date', '')
            qty = data.get('qty', '')
            po_no = data.get('po_no', '')

            if filename.lower().endswith('.xlsx'):
                wb = openpyxl.load_workbook(filepath)
                ws = wb.active
                ws['B7'] = code
                ws['C7'] = delivery_date
                ws['G7'] = code
                wb.save(out_filepath)
                success_count += 1
                processed_files.append(filename)
                
            elif filename.lower().endswith('.csv'):
                with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
                    lines = f.read().splitlines()
                    
                out_lines = []
                for idx, line in enumerate(lines):
                    parts = line.split(',')
                    if len(parts) < 2:
                        parts = parts + [''] * (2 - len(parts))
                        
                    if idx == 5: parts[1] = str(code) 
                    elif idx == 6: parts[1] = str(code) 
                    elif idx == 7: parts[1] = str(qty) 
                    elif idx == 11: parts[1] = str(delivery_date) 
                    elif idx == 13: parts[1] = str(po_no) 
                    
                    out_lines.append(','.join(parts))
                    
                with open(out_filepath, 'w', encoding='utf-8-sig', newline='') as f:
                    f.write('\n'.join(out_lines))
                success_count += 1
                processed_files.append(filename)

        if success_count > 0:
            os.startfile(out_dir)

        return {"success": True, "count": success_count, "out_dir": out_dir, "files": processed_files}

    except Exception as e:
        import traceback
        return JSONResponse({"error": str(e), "trace": traceback.format_exc()}, status_code=500)


@app.get("/")
def index():
    html = """
    <!DOCTYPE html>
    <html lang="zh-TW">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>生產履歷與COA 批次處理系統 (本地直出版)</title>
        <style>
            body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; background-color: #f4f7f6; padding: 20px; }
            .container { max-width: 900px; margin: 0 auto; background: white; padding: 30px; border-radius: 12px; box-shadow: 0 4px 15px rgba(0,0,0,0.1); }
            h1 { color: #2c3e50; text-align: center; }
            .btn-blue { background: #3498db; color: white; padding: 10px 20px; border: none; border-radius: 5px; cursor: pointer; font-size: 16px; margin-bottom: 10px; }
            .btn-blue:hover { background: #2980b9; }
            .btn-green { background: #27ae60; color: white; padding: 15px; border: none; border-radius: 8px; font-size: 1.1em; cursor: pointer; width: 100%; margin-top: 20px; }
            .btn-green:hover { background: #219a52; }
            input[type="text"] { width: 100%; padding: 10px; border: 1px solid #ddd; border-radius: 5px; box-sizing: border-box; margin-bottom: 20px; }
            .modal { display: none; position: fixed; z-index: 1; left: 0; top: 0; width: 100%; height: 100%; background-color: rgba(0,0,0,0.5); }
            .modal-content { background-color: #fff; margin: 5% auto; padding: 20px; border: 1px solid #888; width: 80%; border-radius: 10px; max-height: 80vh; overflow-y: auto; }
            .close { color: #aaa; float: right; font-size: 28px; font-weight: bold; cursor: pointer; }
            table { width: 100%; border-collapse: collapse; margin-top: 15px; }
            th, td { padding: 10px; border-bottom: 1px solid #ddd; text-align: left; }
            th { background-color: #f2f2f2; position: sticky; top: 0; }
            .selected-list { margin-top: 20px; background: #e8f4f8; padding: 15px; border-radius: 5px; }
            #status { margin-top: 20px; font-weight: bold; text-align: center; }
        </style>
    </head>
    <body>
        <div class="container">
            <h1>📄 生產履歷與COA 批次處理系統 (本地直出)</h1>
            
            <label style="font-weight:bold;">1. 來源資料夾路徑 (包含要處理的 xlsx/csv)：</label>
            <input type="text" id="sourceDir" value="D:\\GOOGLE ANGET\\N系列報表產生\\NSE-1106A" placeholder="例如 D:\\GOOGLE ANGET\\N系列報表產生\\NSE-1106A">
            
            <label style="font-weight:bold;">2. 選擇出貨資料：</label><br>
            <button class="btn-blue" onclick="loadShipping()">📋 從 Excel 匯入排程 (選擇出貨表資料)</button>
            
            <div class="selected-list" id="selectedList" style="display:none;">
                <h3>已選擇的批號：</h3>
                <ul id="selectedUl"></ul>
            </div>

            <button class="btn-green" id="processBtn" onclick="startProcessing()" style="display:none;">🚀 開始處理並產生至本地資料夾</button>
            <div id="status"></div>
        </div>

        <!-- Modal -->
        <div id="shipModal" class="modal">
            <div class="modal-content">
                <span class="close" onclick="document.getElementById('shipModal').style.display='none'">&times;</span>
                <h2>Excel 匯入筆數與範圍選擇</h2>
                <div style="margin-bottom: 10px;">
                    <button class="btn-blue" onclick="checkAll(true)">全選</button>
                    <button class="btn-blue" onclick="checkAll(false)">全不選</button>
                    <button class="btn-green" style="width:auto; margin-top:0; padding:10px 20px;" onclick="confirmSelection()">✅ 確認勾選並匯入系統</button>
                </div>
                <table>
                    <thead>
                        <tr>
                            <th>選取</th>
                            <th>到貨日期</th>
                            <th>批號</th>
                            <th>到貨地點</th>
                            <th>採購單號</th>
                            <th>數量</th>
                            <th>Sheet</th>
                        </tr>
                    </thead>
                    <tbody id="shipTableBody"></tbody>
                </table>
            </div>
        </div>

        <script>
            let allShipments = [];
            let selectedBatches = [];

            async function loadShipping() {
                const status = document.getElementById('status');
                status.innerText = '讀取中...';
                try {
                    const res = await fetch('/api/shipping_list');
                    const json = await res.json();
                    if (!res.ok) throw new Error(json.error);
                    
                    allShipments = json.data;
                    const tbody = document.getElementById('shipTableBody');
                    tbody.innerHTML = '';
                    
                    allShipments.forEach((item, index) => {
                        const tr = document.createElement('tr');
                        tr.innerHTML = `
                            <td><input type="checkbox" class="ship-check" value="${index}"></td>
                            <td>${item.delivery_date}</td>
                            <td>${item.batch_no}</td>
                            <td>${item.location}</td>
                            <td>${item.po_no}</td>
                            <td>${item.qty}</td>
                            <td>${item.sheet}</td>
                        `;
                        tbody.appendChild(tr);
                    });
                    
                    document.getElementById('shipModal').style.display = 'block';
                    status.innerText = '';
                } catch (e) {
                    status.innerText = '讀取失敗: ' + e.message;
                    status.style.color = 'red';
                }
            }

            function checkAll(check) {
                document.querySelectorAll('.ship-check').forEach(cb => cb.checked = check);
            }

            function confirmSelection() {
                selectedBatches = [];
                document.querySelectorAll('.ship-check:checked').forEach(cb => {
                    selectedBatches.push(allShipments[cb.value]);
                });
                
                document.getElementById('shipModal').style.display = 'none';
                
                const ul = document.getElementById('selectedUl');
                ul.innerHTML = '';
                selectedBatches.forEach(item => {
                    const li = document.createElement('li');
                    li.innerText = `批號: ${item.batch_no} | 地點: ${item.location} | 到貨: ${item.delivery_date}`;
                    ul.appendChild(li);
                });
                
                document.getElementById('selectedList').style.display = 'block';
                document.getElementById('processBtn').style.display = 'block';
            }

            async function startProcessing() {
                const dir = document.getElementById('sourceDir').value.trim();
                if (!dir) return alert('請輸入來源資料夾路徑');
                if (selectedBatches.length === 0) return alert('請先選擇出貨資料');
                
                const status = document.getElementById('status');
                status.innerText = '處理中，請稍候...';
                status.style.color = 'black';
                
                try {
                    const res = await fetch('/api/process_local', {
                        method: 'POST',
                        headers: { 'Content-Type': 'application/json' },
                        body: JSON.stringify({
                            source_folder: dir,
                            selected_batches: selectedBatches
                        })
                    });
                    const json = await res.json();
                    if (!res.ok) throw new Error(json.error);
                    
                    if (json.count > 0) {
                        status.innerText = `✅ 成功處理 ${json.count} 個檔案！資料夾已自動開啟：\n${json.out_dir}`;
                        status.style.color = 'green';
                    } else {
                        status.innerText = `⚠️ 處理完成，但沒有找到符合批號的檔案。`;
                        status.style.color = 'orange';
                    }
                } catch (e) {
                    status.innerText = '處理失敗: ' + e.message;
                    status.style.color = 'red';
                }
            }
        </script>
    </body>
    </html>
    """
    return HTMLResponse(html)

if __name__ == "__main__":
    ip = get_local_ip()
    port = find_available_port(8002)
    print(f"==================================================")
    print(f"Server is running (生產履歷與COA_系統)")
    print(f"Local URL: http://localhost:{port}")
    print(f"Network URL: http://{ip}:{port}")
    print(f"==================================================")
    uvicorn.run(app, host="0.0.0.0", port=port, log_level="error")
