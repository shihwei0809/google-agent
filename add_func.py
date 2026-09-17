import codecs
path = r'd:\GOOGLE ANGET\三合一單網頁架機伺服器\server.py'
with codecs.open(path, 'r', 'utf-8-sig') as f:
    text = f.read()

func = """
def _extract_batch_from_coa_bytes(content, ext):
    import io
    import csv
    try:
        f = io.StringIO(content.decode('utf-8-sig', errors='ignore'))
        reader = csv.reader(f)
        for row in reader:
            if len(row) >= 2 and 'RawLotId' in str(row[0]):
                return str(row[1]).strip()
    except:
        pass
    try:
        f = io.StringIO(content.decode('cp950', errors='ignore'))
        reader = csv.reader(f)
        for row in reader:
            if len(row) >= 2 and 'RawLotId' in str(row[0]):
                return str(row[1]).strip()
    except:
        pass
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(min_row=1, max_row=50, min_col=1, max_col=2):
                if row[0].value and 'RawLotId' in str(row[0].value) and row[1].value:
                    return str(row[1].value).strip()
    except:
        pass
    return None

"""

if 'def _extract_batch_from_coa_bytes' not in text:
    text = text.replace('def extract_tank_from_batch(batch_str: str) -> str:', func + 'def extract_tank_from_batch(batch_str: str) -> str:')

with codecs.open(path, 'w', 'utf-8-sig') as f:
    f.write(text)
