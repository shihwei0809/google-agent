import os

def patch_factory_validation(path):
    if not os.path.exists(path): return
    with open(path, 'r', encoding='utf-8') as f:
        content = f.read()

    old_block = '''        # 比對排程批號 vs 生產履歷批號
        try:
            import openpyxl as _opxl
            lorry_batches = set()
            if hasattr(self, "imported_lorry_files"):
                for l_file in self.imported_lorry_files:
                    try:
                        _wb = _opxl.load_workbook(l_file, data_only=True)
                        _ws = _wb.active
                        for _r in range(7, _ws.max_row + 1):
                            _val = str(_ws.cell(row=_r, column=1).value or "").strip().upper()
                            if _val:
                                lorry_batches.add(_val)
                        _wb.close()
                    except:
                        pass

            table_batches = [row["batch_var"].get().strip().upper() for row in self.entries if row["batch_var"].get().strip()]

            missing = [b for b in table_batches if b not in lorry_batches]
            found   = [b for b in table_batches if b in lorry_batches]'''

    new_block = '''        # 比對排程批號 vs 生產履歷批號
        try:
            import openpyxl as _opxl
            lorry_batches = set()
            lorry_factory_info = {}
            if hasattr(self, "imported_lorry_files"):
                for l_file in self.imported_lorry_files:
                    try:
                        _wb = _opxl.load_workbook(l_file, data_only=True)
                        _ws = _wb.active
                        for _r in range(7, _ws.max_row + 1):
                            _val = str(_ws.cell(row=_r, column=1).value or "").strip().upper()
                            if _val:
                                lorry_batches.add(_val)
                                f_code = str(_ws.cell(row=_r, column=2).value or "").strip().upper()
                                lorry_factory_info[_val] = f_code
                        _wb.close()
                    except:
                        pass

            table_entries = []
            for row in self.entries:
                b = row["batch_var"].get().strip().upper()
                if b:
                    table_entries.append((b, row["long_code_var"].get().strip().upper()))

            table_batches = [b for b, lc in table_entries]
            missing = [b for b in table_batches if b not in lorry_batches]
            found   = [b for b in table_batches if b in lorry_batches]
            
            wrong_factory_msgs = []
            for b, lc in table_entries:
                if b in lorry_batches:
                    l_fac = lorry_factory_info.get(b, "")
                    expected_substr = lc[1:5] if len(lc) >= 5 else lc
                    if not l_fac:
                        wrong_factory_msgs.append(f"⚠️ 批號 {b}：廠區空白 (應含 {expected_substr})")
                    elif l_fac not in lc:
                        wrong_factory_msgs.append(f"⚠️ 批號 {b}：廠區錯誤 ({l_fac})，未對齊長代號 ({lc})")'''

    content = content.replace(old_block, new_block)

    old_lines_logic = '''            if not table_batches:
                lines.append("ℹ️ 排程表格尚未輸入批號，無法比對。")
            else:
                for b in found:
                    lines.append(f"✅ {b} 已在 生產履歷 中找到")
                for b in missing:
                    lines.append(f"❌ {b} 在 生產履歷 中找不到，請確認批號是否輸入錯誤！")
            
            title = "生產履歷已載入" if not missing else "⚠️ 生產履歷載入 (有批號不符)"'''

    new_lines_logic = '''            if not table_batches:
                lines.append("ℹ️ 排程表格尚未輸入批號，無法比對。")
            else:
                for b in found:
                    lines.append(f"✅ {b} 已在 生產履歷 中找到")
                if wrong_factory_msgs:
                    lines.append("\\n--- 廠區異常提醒 ---")
                    lines.extend(wrong_factory_msgs)
                    lines.append("--------------------\\n")
                for b in missing:
                    lines.append(f"❌ {b} 在 生產履歷 中找不到，請確認批號是否輸入錯誤！")
            
            title = "生產履歷已載入" if not (missing or wrong_factory_msgs) else "⚠️ 生產履歷載入 (有異常)"'''

    content = content.replace(old_lines_logic, new_lines_logic)

    with open(path, 'w', encoding='utf-8') as f:
        f.write(content)

patch_factory_validation(r'D:\GOOGLE ANGET\三合一單自動產生器\main.py')
patch_factory_validation(r'D:\GOOGLE ANGET\勝一三合一單產生系統\main.py')
