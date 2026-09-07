# -*- coding: utf-8 -*-
import os
import sys

# Ensure UTF-8 output in Windows console
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import win32com.client
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml, OxmlElement

base_dir = os.path.dirname(os.path.abspath(__file__))

def set_cell_background(cell, fill_color):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), fill_color))
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'<w:tcMar {}>'
                      r'<w:top w:w="{}" w:type="dxa"/>'
                      r'<w:left w:w="{}" w:type="dxa"/>'
                      r'<w:bottom w:w="{}" w:type="dxa"/>'
                      r'<w:right w:w="{}" w:type="dxa"/>'
                      r'</w:tcMar>'.format(nsdecls('w'), top, left, bottom, right))
    tcPr.append(tcMar)

def add_callout(doc, text_list, title="【重點提醒】", bg_color="F0FDF4", border_color="16A34A"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'<w:tcBorders {}>'
                          r'<w:top w:val="none"/>'
                          r'<w:left w:val="single" w:sz="36" w:space="0" w:color="{}"/>'
                          r'<w:bottom w:val="none"/>'
                          r'<w:right w:val="none"/>'
                          r'</w:tcBorders>'.format(nsdecls('w'), border_color))
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r_t = p.add_run(title + "\n")
    r_t.font.name = 'Microsoft JhengHei'
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor.from_string(border_color)
    
    for i, item in enumerate(text_list):
        r = p.add_run(item + ("\n" if i < len(text_list) - 1 else ""))
        r.font.name = 'Microsoft JhengHei'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def style_table_header(row, titles, bg_color="1E3A8A"):
    for idx, title in enumerate(titles):
        cell = row.cells[idx]
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.font.name = 'Microsoft JhengHei'
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

def style_table_row(row, values, is_even=False, aligns=None):
    bg = "F8FAFC" if is_even else "FFFFFF"
    for idx, val in enumerate(values):
        cell = row.cells[idx]
        set_cell_background(cell, bg)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        if aligns and idx < len(aligns):
            p.alignment = aligns[idx]
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(str(val))
        run.font.name = 'Microsoft JhengHei'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def create_manual():
    docx_path = os.path.join(base_dir, "三合一單與COA雙重核對系統_操作手冊.docx")
    pdf_path = os.path.join(base_dir, "三合一單與COA雙重核對系統_操作手冊.pdf")
    
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Microsoft JhengHei'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # 1. Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("台積三合一單與 COA 雙重核對系統\n核對邏輯與 GAS 線上標準操作手冊")
    r_title.font.name = 'Microsoft JhengHei'
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("GAS 雲端智慧辨識版 | 雙重瀑布流 OCR | 四道安全卡控防線 | 完整現場與管理指南")
    r_sub.font.name = 'Microsoft JhengHei'
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    add_callout(doc, [
        "• 系統架構：Google Apps Script (GAS) Web App + 響應式手機網頁 (相容 Android / iOS / 電腦端)",
        "• 核心使命：徹底解決台積電槽車出貨人工核對 COA 報告與地磅單據時易眼花、誤判槽號、跑錯廠區之風險",
        "• 雙重防護：Google Cloud Vision 主力 OCR (1000次/月) + Gemini 多模型瀑布流自動備援，服務永不中斷",
        "• 四重卡控：COA 批號、送達地點 8 碼拆分、地磅槽號反干擾提取、地磅畫面批號，4 項全數命中才放行"
    ], title="【系統核心特點速覽】", bg_color="EFF6FF", border_color="2563EB")

    # Section 1
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(12)
    p_h1.paragraph_format.space_after = Pt(6)
    r_h1 = p_h1.add_run("一、 系統核心架構與設計理念")
    r_h1.font.name = 'Microsoft JhengHei'
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    p_body1 = doc.add_paragraph()
    p_body1.paragraph_format.line_spacing = 1.25
    p_body1.paragraph_format.space_after = Pt(6)
    p_body1.add_run(
        "在化學品物流與半導體供應鏈中，台積電槽車出貨作業具備極高的品管標準。每次出車均需核對三合一單 QR Code、"
        "COA (品管檢驗分析報告書) 以及地磅電腦畫面之磅單與送達地點。傳統作業仰賴現場人員肉眼逐字比對，"
        "面臨字體細小、夜晚光線昏暗、車次緊湊及人為疏失等挑戰。\n\n"
        "本系統採用 GAS (Google Apps Script) 雲端架構開發，現場人員只需開啟手機瀏覽器即可作業，"
        "無需安裝任何特定 App。系統整合高精度光學字元辨識 (OCR) 與高容錯比對演算法，"
        "於 2~3 秒內自動完成雙重核對，並自動將檢驗合格數據與高解析照片存證歸檔至雲端硬碟。"
    )

    # OCR Table
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    tbl_ocr = doc.add_table(rows=5, cols=4)
    tbl_ocr.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_header(tbl_ocr.rows[0], ["階層 / 角色", "服務模組", "免費額度與調用策略", "備援與容錯機制"])
    style_table_row(tbl_ocr.rows[1], ["第 1 階 (主力)", "Google Cloud Vision API", "雙金鑰輪替，每月享 2,000 次高精度辨識", "設定 900 次安全切換閥值，達標自動換 Key"], False)
    style_table_row(tbl_ocr.rows[2], ["第 2 階 (備援1)", "Gemini 3.6 Flash", "每日免費額度 20 次，高智能 OCR 語意理解", "Vision API 耗盡或錯誤時，自動無縫秒級接管"], True)
    style_table_row(tbl_ocr.rows[3], ["第 3 階 (備援2)", "Gemini 3.5 Flash", "每日免費額度 20 次，快速視覺辨識備援", "3.6 Flash 滿載時自動順位降級調用"], False)
    style_table_row(tbl_ocr.rows[4], ["第 4 階 (常態備援)", "Gemini Flash-Lite 系列", "每日高達 500 次呼叫額度 (3.5 / 3.1 lite)", "大出貨量保底機制，確保系統 24H 永不卡單"], True)

    # Section 2
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(14)
    p_h2.paragraph_format.space_after = Pt(6)
    r_h2 = p_h2.add_run("二、 核心核對邏輯與防呆演算法剖析")
    r_h2.font.name = 'Microsoft JhengHei'
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    p_body2 = doc.add_paragraph()
    p_body2.paragraph_format.line_spacing = 1.25
    p_body2.paragraph_format.space_after = Pt(6)
    p_body2.add_run(
        "系統核對核心函式為 processFormAndVerify_V10(formObject)。"
        "前端提交包含【QR Code 資訊】、【COA 照片 Base64】及【地磅螢幕照片 Base64】。"
        "伺服器端必須通過嚴格的「四重關卡」同時命中，才判定為成功並寫入 Data 試算表；任一項不符即駁回並精準提示錯誤。"
    )

    # 插入四道檢驗關卡核心決策流程圖 (邏輯圖)
    img_logic_path = os.path.join(base_dir, 'manual_assets', '核對邏輯流程圖.png')
    if os.path.exists(img_logic_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(img_logic_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run("▲ 圖 2-1：台積三合一單與 COA 雙重核對系統四道檢驗關卡核心決策流程圖")
        r_cap.font.name = 'Microsoft JhengHei'
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    tbl_logic = doc.add_table(rows=5, cols=4)
    tbl_logic.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_header(tbl_logic.rows[0], ["核對關卡", "比對標的與函式", "核心判讀演算法", "防呆與容錯設計"])
    style_table_row(tbl_logic.rows[1], [
        "關卡 1：\nCOA 批號核對",
        "COA 照片 OCR文字\nvs\nQR 目標批號\n(advancedFuzzyCheck)",
        "1. 長度>=11碼時自動截取 1~11 碼主批號。\n2. 去除所有空白與橫線。\n3. 字元易混淆替換 (3↔5, 8↔B, O↔0, I↔1, Z↔2, S↔5)。",
        "防止檢驗人員拿錯隔壁槽或前一車次的 COA 報告；相容紙張陰影與摺痕辨識瑕疵。"
    ], False)
    style_table_row(tbl_logic.rows[2], [
        "關卡 2：\n送達地點核對",
        "地磅照片 OCR文字\nvs\n目標送達地點\n(smartLocationCheck)",
        "1. 自動剔除開頭 'E' (例如 EF180183B → F180183B)。\n2. 核心 8 碼拆分：前4碼廠別(F180) + 後4碼廠區(183B)。\n3. OCR 內文必須同時包含前4碼與後4碼。\n4. 反光 2 與 3 雙向容錯。",
        "徹底避免台積電不同廠別（如 F18 與 F14）或廠區代碼混淆跑錯門口。"
    ], True)
    style_table_row(tbl_logic.rows[3], [
        "關卡 3：\n地磅槽號核對",
        "地磅照片 OCR文字\nvs\n目標槽號\n(advancedFuzzyCheck)",
        "★ 獨家批號反干擾演算法：\n槽號(如 18P3B、E44) 字元短，易被地磅長批號中的連續字母數字誤判。\n系統比對前先自地磅 OCR 中將目標批號抽離挖空，再比對槽號。",
        "杜絕長批號內部剛好含有槽號字元時產生的「假陽性」命中，百分之百精確。"
    ], False)
    style_table_row(tbl_logic.rows[4], [
        "關卡 4：\n地磅畫面批號",
        "地磅照片 OCR文字\nvs\n目標出貨批號\n(advancedFuzzyCheck)",
        "地磅系統當前掛入之磅單車次，其顯示之生產批號必須與 QR Code 批號一致。",
        "防止地磅員刷錯單、掛錯車次或槽車提早過磅導致資料錯置。"
    ], True)

    p_subh = doc.add_paragraph()
    p_subh.paragraph_format.space_before = Pt(8)
    p_subh.paragraph_format.space_after = Pt(4)
    r_subh = p_subh.add_run("★ 單號智慧抽取與雙月自動滾動建檔")
    r_subh.font.name = 'Microsoft JhengHei'
    r_subh.font.size = Pt(11.5)
    r_subh.font.bold = True
    r_subh.font.color.rgb = RGBColor(0x00, 0x20, 0x60)

    p_body2b = doc.add_paragraph()
    p_body2b.paragraph_format.line_spacing = 1.2
    p_body2b.paragraph_format.space_after = Pt(6)
    p_body2b.add_run(
        "1. 單號自動正則抽取：\n"
        "   - 來源單號：正則表達式 \\b(ESXM1[A-Z0-9\\-]+) 自動捕捉。\n"
        "   - 磅單編號：正則表達式 \\b(ESXM2[A-Z0-9\\-]+) 或 (?:單據編號|單據)[^\\w\\d]*([A-Z0-9\\-]+) 自動捕捉。\n"
        "2. 雙月自動分頁歸檔 (getTargetSheetName)：\n"
        "   - 系統依當月月份自動建立如「Data_2026_09-10」分頁，每兩個月滾動開啟新工作表，避免破萬筆資料拖慢 Google 試算表查詢速度。\n"
        "3. 雲端硬碟高清直存：\n"
        "   - 核對成功後，COA 照片與地磅照片自動轉換為二進位 Blob 存入指定 Google Drive 資料夾，並將直連檢視 URL 寫入試算表。"
    )

    # Section 3
    p_h3 = doc.add_paragraph()
    p_h3.paragraph_format.space_before = Pt(14)
    p_h3.paragraph_format.space_after = Pt(6)
    r_h3 = p_h3.add_run("三、 現場人員線上標準作業手冊 (SOP)")
    r_h3.font.name = 'Microsoft JhengHei'
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    steps = [
        ("步驟 1：開啟系統並載入出貨資料", [
            "• 使用手機 Chrome、Safari 或平板開啟 GAS Web App 連結。",
            "• 方式 A (推薦 - 掃描 QR)：點擊頂部藍色【掃描 QR】按鈕，鏡頭對準三合一單上方 QR Code，1 秒內自動填入料號、槽號、批號、供應商與地點。",
            "• 方式 B (備用 - AI 讀單)：若 QR Code 污損，點擊紫色【AI 讀單】對準三合一單紙本標籤拍照，AI 自動以正則解析並帶入表單。"
        ]),
        ("步驟 2：拍攝 COA 檢驗報告照片", [
            "• 點擊『📷 1. 拍攝 COA 報告照片』虛線方框。",
            "• 確保 COA 報告上方之『批號 (Batch No.)』與檢驗判定區域清晰可見，按下拍照確認。",
            "• 系統自動將照片預覽轉為灰階高對比，強化 OCR 文字辨識度。"
        ]),
        ("步驟 3：拍攝 地磅電腦螢幕照片", [
            "• 點擊『📷 2. 拍攝 地磅照片』虛線方框。",
            "• 鏡頭對準地磅電腦畫面，確保包含：送達地點代碼、槽號、批號、ESXM 來源單號/磅單單號。",
            "• 注意：避免日光燈或窗戶在螢幕上的嚴重反光，正對螢幕拍攝最佳。"
        ]),
        ("步驟 4：送出雙重核對與結果處置", [
            "• 點擊底部綠色【🚀 送出雙重核對】按鈕，系統背景呼叫 Vision/Gemini OCR 進行 4 關卡交叉驗證（約 2~3 秒）。",
            "• 【核對成功處置】：跳出綠色彈窗『✅ 完美！全數核對成功』，顯示擷取之 ESXM1/ESXM2 單號，資料自動存入 Google 試算表，表單重置，可接續作業。",
            "• 【核對失敗處置】：彈出紅色警示窗，標明失敗原因：",
            "   - 找不到批號：檢查 COA 是否拿錯，或重拍避開反光陰影。",
            "   - 地點不符：核對地磅掛入地點與派車排程是否相符。",
            "   - 找不到槽號：確認地磅螢幕槽號是否有遮擋或輸入錯誤。",
            "   - 畫面批號錯誤：核對地磅作業車次是否正確。"
        ])
    ]

    for s_title, s_bullets in steps:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.space_before = Pt(6)
        p_st.paragraph_format.space_after = Pt(2)
        r_st = p_st.add_run(s_title)
        r_st.font.name = 'Microsoft JhengHei'
        r_st.font.bold = True
        r_st.font.size = Pt(11)
        r_st.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        for b in s_bullets:
            p_sb = doc.add_paragraph(style='List Bullet')
            p_sb.paragraph_format.space_after = Pt(2)
            r_sb = p_sb.add_run(b)
            r_sb.font.name = 'Microsoft JhengHei'
            r_sb.font.size = Pt(10)

    # Section 4
    p_h4 = doc.add_paragraph()
    p_h4.paragraph_format.space_before = Pt(14)
    p_h4.paragraph_format.space_after = Pt(6)
    r_h4 = p_h4.add_run("四、 歷史紀錄查詢與報表匯出")
    r_h4.font.name = 'Microsoft JhengHei'
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    p_body4 = doc.add_paragraph()
    p_body4.paragraph_format.line_spacing = 1.25
    p_body4.paragraph_format.space_after = Pt(6)
    p_body4.add_run(
        "1. 即時查詢介面：點擊首頁右上角【查詢紀錄】按鈕進入。\n"
        "2. 關鍵字快搜：輸入批號、料號、來源單號或磅單編號，系統自動於所有歷史 Data 工作表跨分頁檢索。\n"
        "3. 日期區間過濾：支援設定起始日期與結束日期，快速鎖定當日或特定出貨週期。\n"
        "4. 高清佐證調閱：表格內每筆紀錄提供【COA照片】與【地點照片】連結，點擊即可開啟 Google Drive 原圖查驗。\n"
        "5. 一鍵匯出 CSV / Excel：點擊【📥 匯出 Excel (CSV)】，系統自動以 UTF-8 BOM 格式產生報表，Excel 開啟文字清晰無亂碼。"
    )

    # Section 5
    p_h5 = doc.add_paragraph()
    p_h5.paragraph_format.space_before = Pt(14)
    p_h5.paragraph_format.space_after = Pt(6)
    r_h5 = p_h5.add_run("五、 管理者後台維護與疑難排解 (FAQ)")
    r_h5.font.name = 'Microsoft JhengHei'
    r_h5.font.size = Pt(14)
    r_h5.font.bold = True
    r_h5.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    tbl_admin = doc.add_table(rows=5, cols=3)
    tbl_admin.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_table_header(tbl_admin.rows[0], ["後台分頁 / 項目", "功能與配置說明", "日常巡檢重點"])
    style_table_row(tbl_admin.rows[1], [
        "『設定』分頁\n(getSystemConfig)",
        "動態配置金鑰清單：\n- VISION_API_KEY (多把輪替)\n- GEMINI_API_KEY (多把備援)\n- GEMINI_MODEL (模型優先序)\n- USAGE_LIMIT (切換閥值，預設900)",
        "每季檢查 Vision API 金鑰帳號是否正常；可直接在試算表抽換金鑰，免重新部署 GAS。"
    ], False)
    style_table_row(tbl_admin.rows[2], [
        "『API_Log』分頁\n(logApiUsage)",
        "記錄每一次辨識呼叫之：\n時間、API 類型、模型名稱、Key 遮罩代碼、執行狀態 (成功/失敗)、報錯原因。",
        "若發現大量失敗紀錄，檢查是否有網路斷線或 Google Cloud 帳號配額不足。"
    ], True)
    style_table_row(tbl_admin.rows[3], [
        "雲端硬碟封存目錄\n(BACKUP_FOLDER_ID)",
        "指定儲存所有出貨現場拍照圖片之 Google Drive 資料夾 ID。",
        "定期檢視 Google Drive 空間容量；若更換資料夾需同步修改 GS 頂部常數。"
    ], False)
    style_table_row(tbl_admin.rows[4], [
        "雙月分頁管理\n(Data_YYYY_MM-MM)",
        "系統全自動雙月建檔，第 1 欄至第 13 欄保存完整結構化數據與單號。",
        "歷史分頁可直接匯出封存，確保生產數據具備可追溯性。"
    ], True)

    add_callout(doc, [
        "• Q1：現場拍照後提示『辨識逾時』或『網路錯誤』？",
        "  解法：現場 Wi-Fi 或 4G 訊號若較弱，照片上傳需稍待 2 秒；系統已內建多模型瀑布流，會自動重試備援模型。",
        "• Q2：地磅照片地點明明正確，卻判定『地點不符』？",
        "  解法：檢查地磅螢幕是否有光斑剛好反光在廠別代號（例如 F18）上，稍微調整角度避開反光即可秒過。",
        "• Q3：Vision API 每月額度何時重置？",
        "  解法：Google Cloud 每月 1 號 00:00 自動歸零；系統亦自動於每月首次呼叫時重設 USAGE_COUNT 計數器。"
    ], title="【常見問題與故障排除 (FAQ)】", bg_color="FFFBEB", border_color="D97706")

    # Save DOCX
    doc.save(docx_path)
    print(f"✅ DOCX 已成功產出: {docx_path}")

    # Convert to PDF
    try:
        word = win32com.client.Dispatch('Word.Application')
        word.Visible = False
        doc_obj = word.Documents.Open(docx_path)
        # wdFormatPDF = 17
        doc_obj.SaveAs(pdf_path, FileFormat=17)
        doc_obj.Close()
        word.Quit()
        print(f"✅ PDF 已成功轉換: {pdf_path}")
    except Exception as e:
        print(f"⚠️ PDF 轉換異常: {e}")

if __name__ == '__main__':
    create_manual()
