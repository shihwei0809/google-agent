"""
勝一化工 - 槽車出貨排程管理系統 簡報生成腳本
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm
import os

OUTPUT_FILE = os.path.join(os.path.dirname(__file__), '勝一槽車排程系統_上級報告.pptx')

# ── 色彩主題 ──────────────────────────────
C_DARK_BG   = RGBColor(0x0D, 0x1B, 0x2A)   # 深藍黑
C_ACCENT    = RGBColor(0x00, 0xB4, 0xD8)   # 亮藍
C_ACCENT2   = RGBColor(0x48, 0xCA, 0xE4)   # 淺藍
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT_GR  = RGBColor(0xE0, 0xF7, 0xFA)
C_GOLD      = RGBColor(0xFF, 0xD1, 0x66)
C_GREEN     = RGBColor(0x2E, 0xCC, 0x71)
C_RED       = RGBColor(0xE7, 0x4C, 0x3C)
C_GRAY      = RGBColor(0xAA, 0xBB, 0xCC)
C_CARD_BG   = RGBColor(0x1A, 0x2E, 0x40)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank_layout = prs.slide_layouts[6]  # completely blank


def add_slide():
    return prs.slides.add_slide(blank_layout)


def bg(slide, color=C_DARK_BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def box(slide, l, t, w, h, color, alpha=None):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def txt(slide, text, l, t, w, h, size=18, bold=False, color=C_WHITE,
        align=PP_ALIGN.LEFT, wrap=True, italic=False):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return txb


def divider(slide, t, color=C_ACCENT):
    box(slide, 0.6, t, 12.13, 0.04, color)


def badge(slide, text, l, t, color=C_ACCENT, text_color=C_WHITE):
    b = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(1.6), Inches(0.38))
    b.fill.solid()
    b.fill.fore_color.rgb = color
    b.line.fill.background()
    tf = b.text_frame
    tf.word_wrap = False
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(11)
    run.font.bold  = True
    run.font.color.rgb = text_color


def card(slide, l, t, w, h, title, body_lines, title_color=C_ACCENT, icon=''):
    box(slide, l, t, w, h, C_CARD_BG)
    # top accent bar
    box(slide, l, t, w, 0.06, title_color)
    txt(slide, f'{icon}  {title}', l+0.1, t+0.1, w-0.2, 0.4,
        size=13, bold=True, color=title_color)
    body = '\n'.join(body_lines)
    txt(slide, body, l+0.15, t+0.55, w-0.3, h-0.65,
        size=11, color=C_WHITE)


# ═══════════════════════════════════════════
# SLIDE 1 – 封面
# ═══════════════════════════════════════════
s = add_slide()
bg(s)
# gradient accent bar left
box(s, 0, 0, 0.12, 7.5, C_ACCENT)
box(s, 0.12, 0, 0.06, 7.5, C_ACCENT2)

txt(s, '勝一化工股份有限公司', 0.5, 1.0, 12, 0.6, size=16, color=C_GRAY, align=PP_ALIGN.CENTER)
txt(s, '槽車出貨排程管理系統', 0.5, 1.7, 12, 1.1, size=44, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
txt(s, '即時派工・雙向確認・全平台管理', 0.5, 2.9, 12, 0.6, size=20, color=C_ACCENT, align=PP_ALIGN.CENTER)
divider(s, 3.65)
txt(s, '系統說明報告　2026 年 9 月', 0.5, 3.8, 12, 0.5, size=15, color=C_GRAY, align=PP_ALIGN.CENTER)

# stats row
for i, (val, label) in enumerate([('35', '系統帳號'), ('24/7', '全天候雲端'), ('< 1s', '即時同步'), ('100%', '已讀回執')]):
    lx = 1.8 + i * 2.4
    box(s, lx, 4.6, 1.9, 1.3, C_CARD_BG)
    box(s, lx, 4.6, 1.9, 0.06, C_ACCENT)
    txt(s, val,   lx, 4.7,  1.9, 0.6, size=28, bold=True, color=C_GOLD,  align=PP_ALIGN.CENTER)
    txt(s, label, lx, 5.3, 1.9, 0.4, size=12, color=C_GRAY, align=PP_ALIGN.CENTER)

txt(s, '機密文件 ─ 僅供內部簡報使用', 0.5, 6.9, 12, 0.4, size=11, color=C_GRAY, align=PP_ALIGN.CENTER, italic=True)


# ═══════════════════════════════════════════
# SLIDE 2 – 系統背景與目的
# ═══════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 13.33, 0.08, C_ACCENT)
txt(s, '01  系統背景與建置目的', 0.5, 0.2, 12, 0.55, size=24, bold=True, color=C_WHITE)
divider(s, 0.9)

problems = [
    '📄  傳統紙本排程  →  人工傳遞效率低，資訊易落差',
    '📞  電話確認      →  來回溝通耗時，派工紀錄難追蹤',
    '🔴  無法確認是否已讀  →  技服人員收到排程無法即時確認',
    '📊  資料分散      →  Excel 各自為政，版本混亂',
    '🌐  人員在外      →  無法即時取得最新派工資訊',
]
txt(s, '▌ 導入前的痛點', 0.5, 1.05, 12, 0.4, size=16, bold=True, color=C_GOLD)
for i, p in enumerate(problems):
    txt(s, p, 0.8, 1.55 + i*0.56, 6.8, 0.5, size=13, color=C_WHITE)

txt(s, '▌ 建置目標', 7.5, 1.05, 5.5, 0.4, size=16, bold=True, color=C_ACCENT)
goals = [
    ('🎯', '統一派工平台', '單一入口，電腦/手機皆可操作'),
    ('📱', '即時推播通知', '人員手機秒收到新排程任務'),
    ('✅', '已讀回執機制', '確保每筆排程皆有人員確認'),
    ('☁️', '雲端 24/7 上線', '不依賴公司網路，外部人員皆可連線'),
]
for i, (icon, title, desc) in enumerate(goals):
    lx, ty = 7.5, 1.55 + i*1.05
    box(s, lx, ty, 5.5, 0.9, C_CARD_BG)
    box(s, lx, ty, 0.06, 0.9, C_ACCENT)
    txt(s, f'{icon} {title}', lx+0.15, ty+0.05, 5.2, 0.38, size=13, bold=True, color=C_ACCENT)
    txt(s, desc, lx+0.15, ty+0.45, 5.2, 0.38, size=11, color=C_GRAY)


# ═══════════════════════════════════════════
# SLIDE 3 – 系統架構總覽
# ═══════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 13.33, 0.08, C_ACCENT)
txt(s, '02  系統架構總覽', 0.5, 0.2, 12, 0.55, size=24, bold=True, color=C_WHITE)
divider(s, 0.9)

# 架構圖 (用形狀模擬)
layers = [
    (C_ACCENT,  '前端介面層',    '電腦管理端 (Chrome/Edge)  ╱  技服/司機手機 PWA App'),
    (C_GREEN,   '應用服務層',    'Node.js + Express  ─  REST API / SSE 即時事件流 / Web Push'),
    (C_GOLD,    '資料持久層',    'JSON 資料庫 (orders / users / drivers)  ╱  Excel 班表同步'),
    (RGBColor(0x9B,0x59,0xB6), '部署基礎設施', 'Render.com 雲端 Docker 容器  ╱  UptimeRobot 24/7 保活監控'),
]
for i, (c, title, desc) in enumerate(layers):
    ty = 1.1 + i * 1.35
    box(s, 0.5, ty, 12.3, 1.15, C_CARD_BG)
    box(s, 0.5, ty, 0.08, 1.15, c)
    txt(s, title, 0.8, ty+0.1, 3.0, 0.4, size=14, bold=True, color=c)
    txt(s, desc,  0.8, ty+0.55, 11.5, 0.5, size=12, color=C_WHITE)

# 箭頭文字
for i in range(3):
    txt(s, '↕', 6.5, 2.17 + i * 1.35, 1, 0.3, size=18, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════
# SLIDE 4 – 角色權限說明
# ═══════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 13.33, 0.08, C_ACCENT)
txt(s, '03  角色權限與帳號設計', 0.5, 0.2, 12, 0.55, size=24, bold=True, color=C_WHITE)
divider(s, 0.9)

roles_data = [
    ('👑', 'admin\n系統管理員',     C_RED,   ['管理所有帳號', '上傳/下載任何資料', '完整日誌查閱', '系統設定']),
    ('📊', 'sales\n業務人員',        C_ACCENT,['上傳空白班表 Excel', '查看訂單狀態', '下載彙整 Excel']),
    ('🔧', 'tech_mgr\n技服主管',     C_GOLD,  ['匯入技服充填手排程', '指定技服人員', '查閱已讀狀態', '發送推播通知']),
    ('🚛', 'transporter\n運輸公司',  RGBColor(0x9B,0x59,0xB6), ['上傳司機/車牌資料', '比對排程', '查看派車狀態']),
    ('👷', 'tech_staff\n技服人員',   C_GREEN, ['查看自己的排程', '標記已讀', '查詢歷史排程', '接收推播通知']),
]
cols = [(0.3, 2.5), (2.95, 2.5), (5.6, 2.5), (8.25, 2.5), (10.9, 2.28)]
for i, ((lx, cw), (icon, role, color, perms)) in enumerate(zip(cols, roles_data)):
    box(s, lx, 1.1, cw, 5.7, C_CARD_BG)
    box(s, lx, 1.1, cw, 0.08, color)
    txt(s, icon, lx, 1.2, cw, 0.6, size=22, align=PP_ALIGN.CENTER)
    txt(s, role, lx, 1.85, cw, 0.65, size=12, bold=True, color=color, align=PP_ALIGN.CENTER)
    for j, perm in enumerate(perms):
        txt(s, f'• {perm}', lx+0.1, 2.65 + j*0.62, cw-0.15, 0.55, size=10, color=C_WHITE)

txt(s, f'共 35 個正式帳號　│　所有帳號密碼預設為 123，可由管理員修改', 0.5, 7.1, 12.3, 0.35,
    size=11, color=C_GRAY, align=PP_ALIGN.CENTER, italic=True)


# ═══════════════════════════════════════════
# SLIDE 5 – 操作流程 (主流程)
# ═══════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 13.33, 0.08, C_ACCENT)
txt(s, '04  完整操作流程', 0.5, 0.2, 12, 0.55, size=24, bold=True, color=C_WHITE)
divider(s, 0.9)

steps = [
    ('1', C_ACCENT,  '業務上傳空白班表', '業務登入 → 匯入 Excel 班表\n系統解析訂單資料並儲存'),
    ('2', C_GOLD,    '技服主管指派充填手', '主管登入 → 比對技服 Excel\n確認比對結果後確認匯入'),
    ('3', C_GOLD,    '主管確認推播通知', '系統顯示「哪些人有新任務」\n主管按確定才送出推播通知'),
    ('4', C_GREEN,   '技服人員收到通知', '手機彈出推播 → 點開系統\n查看自己當天/明天派工任務'),
    ('5', C_GREEN,   '人員確認已讀', '點擊「已讀」回傳確認\n電腦端即時同步顯示已讀'),
    ('6', RGBColor(0x9B,0x59,0xB6), '運輸公司補充車輛', '司機/車牌資料比對匯入\n完成完整出貨資訊'),
    ('7', C_ACCENT,  '匯出彙整報表', '下載包含所有資訊的 Excel\n作為出貨/存檔依據'),
]
step_w = 1.76
for i, (num, color, title, desc) in enumerate(steps):
    lx = 0.22 + i * (step_w + 0.08)
    box(s, lx, 1.1, step_w, 5.5, C_CARD_BG)
    box(s, lx, 1.1, step_w, 0.08, color)
    # circle number
    circle = s.shapes.add_shape(9, Inches(lx + step_w/2 - 0.27), Inches(1.2), Inches(0.54), Inches(0.54))
    circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.line.fill.background()
    tf = circle.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = num; run.font.size = Pt(14); run.font.bold = True; run.font.color.rgb = C_DARK_BG

    txt(s, title, lx+0.08, 1.88, step_w-0.12, 0.55, size=12, bold=True, color=color, align=PP_ALIGN.CENTER)
    txt(s, desc,  lx+0.1,  2.55, step_w-0.15, 2.0,  size=10, color=C_WHITE,  align=PP_ALIGN.CENTER)

    if i < 6:
        txt(s, '→', lx + step_w - 0.02, 2.8, 0.25, 0.4, size=14, bold=True, color=C_GRAY, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════
# SLIDE 6 – 即時同步與已讀機制
# ═══════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 13.33, 0.08, C_ACCENT)
txt(s, '05  即時同步與已讀回執機制', 0.5, 0.2, 12, 0.55, size=24, bold=True, color=C_WHITE)
divider(s, 0.9)

# SSE 說明
txt(s, '▌ Server-Sent Events (SSE) 即時推送', 0.5, 1.05, 8, 0.4, size=15, bold=True, color=C_GOLD)
sse_items = [
    '技服人員按下「已讀」後，Server 即時廣播事件給所有在線端',
    '電腦管理端的表格徽章在 < 1 秒內從 🔴 未讀 切換為 🟢 已讀',
    '已讀時間與確認人員同步顯示，無需手動刷新',
    '同時提供 3 秒輪詢備援，確保 SSE 中斷時也能自動恢復同步',
]
for i, item in enumerate(sse_items):
    txt(s, f'✓  {item}', 0.8, 1.55 + i * 0.55, 7.8, 0.48, size=12, color=C_WHITE)

# 已讀保留機制
txt(s, '▌ 重新匯入的智慧已讀保留', 0.5, 3.5, 8, 0.4, size=15, bold=True, color=C_ACCENT)
preserve_items = [
    '同一筆訂單重新匯入時，若負責人員沒有更換 → 已讀狀態保留不動',
    '若負責人員更換 → 自動重置為未讀，確保新負責人收到通知',
    '已讀人員不會重複收到相同任務的推播通知',
]
for i, item in enumerate(preserve_items):
    txt(s, f'✓  {item}', 0.8, 4.0 + i * 0.55, 7.8, 0.48, size=12, color=C_WHITE)

# 右側視覺
box(s, 8.8, 1.05, 4.3, 5.8, C_CARD_BG)
box(s, 8.8, 1.05, 4.3, 0.06, C_ACCENT)
txt(s, '📱 技服人員視角', 8.9, 1.15, 4.1, 0.4, size=13, bold=True, color=C_ACCENT)
phone_lines = [
    ('🔔 推播通知', C_GOLD),
    ('  「新派工任務：林介文」', C_WHITE),
    ('', C_WHITE),
    ('📋 點開排程卡片', C_ACCENT),
    ('  出貨日期 ｜ 到貨時間', C_WHITE),
    ('  客戶地點 ｜ 產品品項', C_WHITE),
    ('', C_WHITE),
    ('✅ 按下「確認已讀」', C_GREEN),
    ('  → 電腦端即時同步', C_GRAY),
    ('  🟢 已讀 by 林聖龍', C_GREEN),
]
for i, (line, c) in enumerate(phone_lines):
    txt(s, line, 9.0, 1.7 + i * 0.48, 3.9, 0.44, size=11, color=c)


# ═══════════════════════════════════════════
# SLIDE 7 – 雲端部署與防休眠機制
# ═══════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 13.33, 0.08, C_ACCENT)
txt(s, '06  雲端部署架構與 24/7 可用性', 0.5, 0.2, 12, 0.55, size=24, bold=True, color=C_WHITE)
divider(s, 0.9)

deploy_cards = [
    ('☁️ Render.com\n雲端主機', C_ACCENT,
     ['免費方案部署 Docker 容器', '自動偵測 GitHub 推送並重新部署',
      '正式網址：google-agent-4gzi.onrender.com', '全球 CDN / Cloudflare 加速']),
    ('👁️ UptimeRobot\n外部監控', C_GREEN,
     ['每 5 分鐘自動 Ping 伺服器', '防止 Render 閒置 15 分鐘自動休眠',
      '偵測到掛掉立即發 Email 通知', '近 30 天可用率 100%']),
    ('🔁 程式內建\n自打心跳', C_GOLD,
     ['server.js 每 10 分鐘對外發送請求', '雙重保險確保伺服器永不休眠',
      '啟動時自動顯示本機 IP 與外網網址', '支援 Port 衝突自動切換']),
    ('🏠 本機備援\n固定 Tunnel', RGBColor(0x9B,0x59,0xB6),
     ['本機 Node.js 同步運行備援', '固定 Tunnel 網址對外服務',
      '斷線自動重連，8 秒後重試', '區網 IP 讓辦公室人員直連']),
]

for i, (title, color, items) in enumerate(deploy_cards):
    lx = 0.4 + i * 3.18
    box(s, lx, 1.1, 3.0, 5.8, C_CARD_BG)
    box(s, lx, 1.1, 3.0, 0.08, color)
    txt(s, title, lx, 1.2, 3.0, 0.65, size=13, bold=True, color=color, align=PP_ALIGN.CENTER)
    for j, item in enumerate(items):
        txt(s, f'• {item}', lx+0.12, 2.0 + j * 0.78, 2.8, 0.7, size=11, color=C_WHITE)


# ═══════════════════════════════════════════
# SLIDE 8 – 效益總結
# ═══════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 13.33, 0.08, C_ACCENT)
txt(s, '07  導入效益與未來規劃', 0.5, 0.2, 12, 0.55, size=24, bold=True, color=C_WHITE)
divider(s, 0.9)

benefits = [
    ('⏱️', '派工效率', '原先人工電話通知\n→ 系統自動推播\n節省 80% 溝通時間'),
    ('✅', '確認透明度', '無法追蹤人員是否收到\n→ 已讀回執即時顯示\n主管隨時掌握狀況'),
    ('📱', '行動化作業', '人員須在辦公室查看\n→ 手機 PWA App\n隨時隨地查閱排程'),
    ('☁️', '資訊安全', '資料散落各 Excel\n→ 統一雲端管理\n帳號權限分級控管'),
]
for i, (icon, title, desc) in enumerate(benefits):
    lx = 0.6 + i * 3.05
    box(s, lx, 1.1, 2.8, 3.2, C_CARD_BG)
    box(s, lx, 1.1, 2.8, 0.08, C_ACCENT)
    txt(s, icon,  lx, 1.25, 2.8, 0.6,  size=24, align=PP_ALIGN.CENTER)
    txt(s, title, lx, 1.9,  2.8, 0.45, size=14, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)
    txt(s, desc,  lx+0.1, 2.45, 2.6, 1.7, size=11, color=C_WHITE, align=PP_ALIGN.CENTER)

txt(s, '▌ 未來規劃', 0.5, 4.5, 5, 0.4, size=15, bold=True, color=C_GOLD)
futures = [
    '📊 報表分析模組   ─ 自動統計各技服人員出勤率與確認速率',
    '🗓️ 行事曆視圖      ─ 月曆化呈現排程，一目瞭然',
    '📷 QR Code 掃碼  ─ 現場掃描快速確認到貨',
    '🔔 逾時未讀自動提醒  ─ 排程前 N 小時若未確認自動催讀',
]
for i, f in enumerate(futures):
    txt(s, f, 0.8, 5.05 + i * 0.5, 12, 0.44, size=12, color=C_WHITE)


# ═══════════════════════════════════════════
# SLIDE 9 – 結語
# ═══════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 0.12, 7.5, C_ACCENT)
box(s, 0.12, 0, 0.06, 7.5, C_ACCENT2)

txt(s, '感謝聆聽', 0.5, 1.8, 12, 1.1, size=52, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
txt(s, '勝一化工 槽車出貨排程管理系統', 0.5, 3.0, 12, 0.6, size=20, color=C_ACCENT, align=PP_ALIGN.CENTER)
divider(s, 3.75)

infos = [
    ('🌐 正式網址', 'https://google-agent-4gzi.onrender.com'),
    ('👤 系統帳號', '35 位人員帳號，預設密碼 123'),
    ('📱 PWA 安裝', '用手機 Chrome 開啟網址 → 加到主畫面'),
    ('🔧 技術支援', 'shihwei（系統管理員）'),
]
for i, (label, val) in enumerate(infos):
    lx = 0.6 + (i % 2) * 6.4
    ty = 4.0 + (i // 2) * 0.8
    txt(s, label, lx, ty, 2.2, 0.45, size=12, bold=True, color=C_GOLD)
    txt(s, val,   lx+2.2, ty, 4.0, 0.45, size=12, color=C_WHITE)

txt(s, '本系統由公司內部自主開發，資料存於雲端伺服器，帳號資訊請妥善保管。', 0.5, 6.9, 12.3, 0.35,
    size=10, color=C_GRAY, align=PP_ALIGN.CENTER, italic=True)


prs.save(OUTPUT_FILE)
print(f'✅ PPT 已生成：{OUTPUT_FILE}')
