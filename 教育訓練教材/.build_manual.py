from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = r"C:\GOOGLE ANGET\教育訓練教材\AI教育訓練平台_操作手冊_使用者版.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)

def set_font(run, size=10.5, bold=False, color="263648"):
    run.font.name = "Microsoft JhengHei"; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)

for stylename,size,color in [("Normal",10.5,"263648"),("Title",26,"12304A"),("Heading 1",17,"12304A"),("Heading 2",13,"167D8D"),("Heading 3",11,"167D8D")]:
    s=doc.styles[stylename]; s.font.name="Microsoft JhengHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)
    if stylename.startswith("Heading"): s.font.bold=True
    if stylename == "Title" and s._element.pPr is not None:
        border=s._element.pPr.find(qn("w:pBdr"))
        if border is not None: s._element.pPr.remove(border)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def p(text="", style=None, boldlead=None):
    para=doc.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        r=para.add_run(boldlead); set_font(r,bold=True,color="12304A")
        r=para.add_run(text[len(boldlead):]); set_font(r)
    else:
        r=para.add_run(text); set_font(r, size=10.5 if not style else None)
    para.paragraph_format.space_after=Pt(4)
    return para

def bullets(items):
    for x in items:
        para=doc.add_paragraph(style="List Bullet"); para.paragraph_format.space_after=Pt(2)
        r=para.add_run(x); set_font(r,10)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style="Table Grid"
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,"12304A"); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: set_font(r,9.5,True,"FFFFFF")
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=val; cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(t.rows)%2==0: shade(cells[i],"F1F6F8")
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after=Pt(1)
                for r in para.runs: set_font(r,9.3,color="263648")
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

# running footer
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run("AI 教育訓練平台｜使用者操作手冊"); set_font(r,8,color="6E7B86")

title=doc.add_paragraph(style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.LEFT
border=title._p.get_or_add_pPr().find(qn("w:pBdr"))
if border is not None: title._p.get_or_add_pPr().remove(border)
r=title.add_run("AI 教育訓練平台操作手冊"); set_font(r,26,True,"12304A")
p("適用對象：一般學習者與教材管理者　｜　依現行平台介面與程式功能編寫")
p("本手冊說明如何啟動平台、瀏覽教材、向 AI 助教提問，以及管理教材。一般學習者可直接瀏覽與提問；上傳、編修及刪除教材須切換管理員模式。")

doc.add_heading("1 使用前準備", level=1)
bullets(["使用 Windows 電腦與可連線至平台主機的瀏覽器。", "本機啟動方式需要專案環境中的 Python、Node.js 及前後端套件；首次使用請由系統管理者完成環境設定。", "AI 問答與教材解析需在 backend/.env 設定有效的 GEMINI_API_KEY，並允許連線至 Gemini API。", "教材檔案存放於後端 materials 資料夾；請依公司資料管理規範上傳與維護內容。"])
doc.add_heading("2 啟動與停止平台", level=1)
doc.add_heading("2.1 啟動", level=2)
p("在平台資料夾雙擊「啟動教育訓練平台.bat」。啟動腳本會檢查環境並開啟後端與前端命令視窗，接著開啟瀏覽器。預設網址為 http://localhost:5173，後端 API 為 http://localhost:8000。")
table(["項目","預設值／說明"], [["前端網址","http://localhost:5173"],["後端 API","http://localhost:8000"],["服務程序","FastAPI 後端與 Vite 前端分別執行"]])
doc.add_heading("2.2 停止", level=2)
p("雙擊「停止教育訓練平台.bat」，依畫面提示停止平台服務。若服務未正常關閉，請由管理者檢查啟動時開啟的命令視窗與程序。")
p("使用區域網路其他電腦時，請由管理者提供主機 IP 與可用網址，並確認防火牆及網路設定允許連線。localhost 僅代表目前這台電腦。")

doc.add_page_break()
doc.add_heading("3 學習者操作", level=1)
doc.add_heading("3.1 瀏覽教材", level=2)
bullets(["開啟平台後，左側顯示「教材列表」。", "點選教材名稱，中央閱讀區載入教材內容。", "教材中的圖片可點擊放大；按 Esc 或點擊背景即可關閉放大檢視。"])
doc.add_heading("3.2 向 AI 助教提問", level=2)
bullets(["在右側「AI 助教」輸入問題，按 Enter 或點擊送出。", "可詢問目前教材重點、步驟說明或名詞解釋；回答會以對話形式顯示。", "回答含圖片時可點擊圖片放大檢視。若回答未引用正確內容，請回看原教材並向講師或主管確認。"])
p("注意：介面會顯示目前對話內容；程式碼未見跨工作階段問答歷史保存功能，因此重新載入或離開後不保證保留對話。AI 回答僅供學習參考，涉及安全、法規、品質或實際作業要求時，應依正式文件及主管指示辦理。")

doc.add_heading("4 教材管理者操作", level=1)
doc.add_heading("4.1 進入管理員模式", level=2)
p("在教材列表右上方點擊「登入」，輸入目前程式設定的預設密碼 admin123。成功後畫面切換為管理員模式，顯示教材上傳與刪除操作。完成維護後再次點擊 Admin 可登出管理員模式。")
p("安全注意：目前前端以固定密碼 admin123 判斷管理員模式，未見帳號、角色伺服器驗證或密碼變更介面。正式使用前應由系統負責人評估並強化存取控制。")
doc.add_heading("4.2 上傳教材", level=2)
bullets(["管理員模式下點擊上傳圖示，選取支援格式：MD、TXT、PDF、DOCX、XLSX、PPTX、MP4、MOV、AVI 或 WEBM。", "等待解析完成。文件型教材會擷取文字並轉為 Markdown；影片會呼叫 Gemini 解析並產生教學內容，處理時間可能較長。", "上傳完成後，教材列表更新並載入新教材。請抽查標題、步驟、圖片和表格是否正確。", "舊版 .doc 不在目前介面列出的支援格式；請先另存為 .docx 再上傳。"])
doc.add_heading("4.3 編修與儲存教材", level=2)
bullets(["選取教材後點擊「編輯教材」，在文字編輯區修改內容。", "可從本機選取圖片，或在編輯區貼上螢幕截圖；依畫面插入圖片後檢查位置與顯示。", "點擊「儲存修改」寫回教材；看到儲存成功訊息後，離開編輯模式並重新檢查內容。"])
doc.add_heading("4.4 刪除教材", level=2)
p("在教材列表點擊該教材旁的刪除圖示，閱讀確認視窗中的檔名後再確認。刪除會移除伺服器端教材檔案；刪除前請先確認該教材不再需要，並依組織備份流程留存必要副本。")
doc.add_heading("4.5 圖片標註", level=2)
p("管理員可在圖片檢視視窗點擊「加框/代號標註」，於圖片加上框線、代號或文字，再儲存。若平台沒有自動同步教材圖片引用，請回到編輯區確認圖片已正確插入並儲存教材。")

doc.add_page_break()
doc.add_heading("5 常見問題排除", level=1)
table(["狀況","建議處理"], [["頁面無法開啟","確認啟動腳本已開啟前後端視窗；本機使用 localhost:5173。"],["教材清單載入失敗","確認後端視窗仍在執行，並檢查 localhost:8000 API 狀態。"],["AI 無回應或解析失敗","請管理者確認 backend/.env 的 GEMINI_API_KEY、網路連線及 Gemini API 可用狀態。"],["上傳格式不支援","確認副檔名在支援清單中；.doc 可先另存為 .docx。"],["教材圖片顯示異常","檢查圖片是否已上傳、Markdown 圖片連結是否正確，並重新儲存教材。"],["管理員登入失敗","確認輸入預設密碼；目前密碼由前端程式固定設定，請系統管理者處理。"]])

doc.add_heading("6 功能範圍與使用限制", level=1)
bullets(["平台目前提供教材瀏覽、AI 對話、管理者上傳/編修/刪除及圖片標註等功能。", "現有程式未顯示員工帳號、課程報名、簽到、測驗、訓練時數、資格到期提醒或 ERP 簽核功能；相關流程仍須使用組織正式系統與表單。", "AI 可能產生不完整或錯誤內容。教材管理者應審閱 AI 轉換結果，學習者應以正式核准教材作業。", "平台存取權限、備份、API 金鑰管理及個人資料處理，應由系統管理者另行依公司規範確認。"])

doc.add_heading("附錄 操作流程速查", level=1)
table(["學習者","教材管理者"], [["開啟平台 → 選教材 → 閱讀 → 提問 → 核對正式教材","登入管理模式 → 上傳/選教材 → 編修 → 儲存 → 抽查結果"]])
doc.save(OUT)
print(OUT)
