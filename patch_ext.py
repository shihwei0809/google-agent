import codecs

path = r'd:\GOOGLE ANGET\勝一三合一單網頁架機伺服器\server.py'
with codecs.open(path, 'r', 'utf-8-sig') as f:
    lines = f.readlines()

new_lines = []
in_func = False
for line in lines:
    if 'def _extract_batch_from_coa_bytes(content, ext):' in line:
        in_func = True
        new_lines.append("""def _extract_batch_from_coa_bytes(content, ext):
    import io
    import csv
    # Try CSV first
    try:
        f = io.StringIO(content.decode('utf-8-sig', errors='ignore'))
        reader = csv.reader(f)
        for row in reader:
            if len(row) >= 2 and 'RawLotId' in str(row[0]):
                return str(row[1]).strip()
    except:
        pass
    try:
        f = io.StringIO(content.decode('cp950', errors='ignore'))
        reader = csv.reader(f)
        for row in reader:
            if len(row) >= 2 and 'RawLotId' in str(row[0]):
                return str(row[1]).strip()
    except:
        pass
    # Try Excel fallback
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(min_row=1, max_row=50, min_col=1, max_col=2):
                if row[0].value and 'RawLotId' in str(row[0].value) and row[1].value:
                    return str(row[1].value).strip()
    except:
        pass
    return None
""")
    elif in_func:
        if 'return None' in line:
            in_func = False
    else:
        new_lines.append(line)

with codecs.open(path, 'w', 'utf-8-sig') as f:
    f.writelines(new_lines)
