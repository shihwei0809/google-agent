import codecs
for path in [r'd:\GOOGLE ANGET\三合一單自動產生器\main.py', r'd:\GOOGLE ANGET\勝一三合一單產生系統\main.py']:
    with codecs.open(path, 'r', 'utf-8-sig') as f:
        lines = f.readlines()
    new_lines = []
    in_func = False
    for line in lines:
        if 'def _extract_batch_from_coa(self, file_path):' in line:
            in_func = True
            new_lines.append("""    def _extract_batch_from_coa(self, file_path):
        import csv
        try:
            with open(file_path, 'r', encoding='utf-8-sig', errors='ignore') as f:
                reader = csv.reader(f)
                for row in reader:
                    if len(row) >= 2 and 'RawLotId' in str(row[0]):
                        return str(row[1]).strip()
        except:
            pass
        try:
            with open(file_path, 'r', encoding='cp950', errors='ignore') as f:
                reader = csv.reader(f)
                for row in reader:
                    if len(row) >= 2 and 'RawLotId' in str(row[0]):
                        return str(row[1]).strip()
        except:
            pass
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows(min_row=1, max_row=50, min_col=1, max_col=2):
                    if row[0].value and 'RawLotId' in str(row[0].value) and row[1].value:
                        return str(row[1].value).strip()
        except:
            pass
        return None
""")
        elif in_func:
            if 'return None' in line:
                in_func = False
        else:
            new_lines.append(line)
    with codecs.open(path, 'w', 'utf-8-sig') as f:
        f.writelines(new_lines)
