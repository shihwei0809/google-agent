import os

def patch_main_py(path):
    if not os.path.exists(path):
        return
    with open(path, 'r', encoding='utf-8') as f:
        content = f.read()

    # The block starts at:
    #         # 比對排程批號 vs 生產履歷批號
    #         try:
    #             import openpyxl as _opxl
    #             _wb = _opxl.load_workbook(filepath, data_only=True)
    #             _ws = _wb.active
    #             lorry_batches = set()
    #             for _r in range(7, _ws.max_row + 1):
    #                 _val = str(_ws.cell(row=_r, column=1).value or "").strip().upper()
    #                 if _val:
    #                     lorry_batches.add(_val)

    old_block = '''        # 比對排程批號 vs 生產履歷批號
        try:
            import openpyxl as _opxl
            _wb = _opxl.load_workbook(filepath, data_only=True)
            _ws = _wb.active
            lorry_batches = set()
            for _r in range(7, _ws.max_row + 1):
                _val = str(_ws.cell(row=_r, column=1).value or "").strip().upper()
                if _val:
                    lorry_batches.add(_val)'''
    
    new_block = '''        # 比對排程批號 vs 生產履歷批號
        try:
            import openpyxl as _opxl
            lorry_batches = set()
            if hasattr(self, "imported_lorry_files"):
                for l_file in self.imported_lorry_files:
                    try:
                        _wb = _opxl.load_workbook(l_file, data_only=True)
                        _ws = _wb.active
                        for _r in range(7, _ws.max_row + 1):
                            _val = str(_ws.cell(row=_r, column=1).value or "").strip().upper()
                            if _val:
                                lorry_batches.add(_val)
                        _wb.close()
                    except:
                        pass'''
                        
    content = content.replace(old_block, new_block)
    
    # Next, we fix `fname` references.
    # old_fname_1 = f"已成功載入生產履歷檔案：\n{fname}\n\n已為您自動勾選【產生單列生產履歷】！\n"
    # old_fname_2 = f"已成功載入生產履歷檔案：\n{fname}\n\n已為您自動勾選【產生單列生產履歷】！\n稍後點擊【開始批次產生】時，系統會自動比對每筆排程批號並單列輸出。"
    
    # To fix this safely, we will just insert `fname = ...` right before `missing = ...` or `lines = ...`
    # Let's do it by inserting `fname` calculation before `lines = [f"已成功...`
    old_lines = '''            lines = [f"已成功載入生產履歷檔案：\\n{fname}\\n\\n已為您自動勾選【產生單列生產履歷】！\\n"]'''
    new_lines = '''            fname = f"已選 {len(self.imported_lorry_files)} 份檔案" if len(self.imported_lorry_files) > 1 else os.path.basename(self.imported_lorry_files[0]) if self.imported_lorry_files else ""
            lines = [f"已成功載入生產履歷檔案：\\n{fname}\\n\\n已為您自動勾選【產生單列生產履歷】！\\n"]'''
    content = content.replace(old_lines, new_lines)
    
    old_except = '''        except Exception as _e:
            messagebox.showinfo(
                "生產履歷已載入", 
                f"已成功載入生產履歷檔案：\\n{fname}\\n\\n已為您自動勾選【產生單列生產履歷】！\\n稍後點擊【開始批次產生】時，系統會自動比對每筆排程批號並單列輸出。"
            )'''
    new_except = '''        except Exception as _e:
            fname = f"已選 {len(self.imported_lorry_files)} 份檔案" if len(self.imported_lorry_files) > 1 else os.path.basename(self.imported_lorry_files[0]) if self.imported_lorry_files else ""
            messagebox.showinfo(
                "生產履歷已載入", 
                f"已成功載入生產履歷檔案：\\n{fname}\\n\\n已為您自動勾選【產生單列生產履歷】！\\n稍後點擊【開始批次產生】時，系統會自動比對每筆排程批號並單列輸出。"
            )'''
    content = content.replace(old_except, new_except)

    with open(path, 'w', encoding='utf-8') as f:
        f.write(content)

patch_main_py(r'D:\GOOGLE ANGET\三合一單自動產生器\main.py')
patch_main_py(r'D:\GOOGLE ANGET\勝一三合一單產生系統\main.py')
print("Patched both main.py files successfully!")
